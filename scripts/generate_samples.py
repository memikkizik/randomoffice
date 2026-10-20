"""Generate realistic Office documents for local API, parser, and roundtrip testing."""

from argparse import ArgumentParser
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def generate_docx(output_dir: Path) -> Path:
    """Create a realistic DOCX with heading, formatted text, Turkish characters, list, and table."""
    path = output_dir / "sample.docx"
    document = Document()
    document.add_heading("RandomOffice Örnek Türkçe Belge", level=1)

    p1 = document.add_paragraph("Bu belge, ayrıştırıcı ve API testleri için oluşturuldu: ")
    r_bold = p1.add_run("Kalın metin")
    r_bold.bold = True
    p1.add_run(" ve ")
    r_italic = p1.add_run("İtalik metin (Şemsiye, Çağlayan, İğne, Ömür, Çözüm).")
    r_italic.italic = True

    document.add_paragraph("Numaralandırılmış Madde 1: İlk adım", style="List Number")
    document.add_paragraph("Numaralandırılmış Madde 2: İkinci adım", style="List Number")

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Alan / Başlık"
    table.cell(0, 1).text = "Değer"
    table.cell(1, 0).text = "Durum"
    table.cell(1, 1).text = "Hazır (Kusursuz)"
    document.save(path)
    return path


def generate_xlsx(output_dir: Path) -> Path:
    """Create a realistic XLSX with multiple sheets, formula, AA column, zero, text ID, and currency."""
    path = output_dir / "sample.xlsx"
    workbook = Workbook()

    # Sheet 1: Özet
    sheet1 = workbook.active
    sheet1.title = "Özet"
    sheet1["A1"] = "RandomOffice Finansal Tablo"
    sheet1["A2"] = "Kimlik No"
    sheet1["B2"] = "Açıklama"
    sheet1["C2"] = "Miktar"
    sheet1["D2"] = "Bakiye"

    # String ID with leading zeros
    sheet1["A3"] = "00123"
    sheet1["B3"] = "Açılış Bakiyesi"
    sheet1["C3"] = 100
    sheet1["D3"] = 1500.50
    sheet1["D3"].number_format = '"₺" #,##0.00'

    # Numeric zero
    sheet1["A4"] = "00124"
    sheet1["B4"] = "Sıfır Bakiye Kalemi"
    sheet1["C4"] = 0
    sheet1["D4"] = 0.00
    sheet1["D4"].number_format = '"₺" #,##0.00'

    # Item 3
    sheet1["A5"] = "00125"
    sheet1["B5"] = "Ek Masraf"
    sheet1["C5"] = 250
    sheet1["D5"] = 250.00
    sheet1["D5"].number_format = '"₺" #,##0.00'

    # Formula row
    sheet1["B6"] = "Toplam Miktar:"
    sheet1["C6"] = "=SUM(C3:C5)"

    # Wide column (AA = column 27)
    sheet1["AA1"] = "Geniş Sütun Verisi"

    # Sheet 2: Detaylar
    sheet2 = workbook.create_sheet(title="Detaylar")
    sheet2["A1"] = "İkinci Sayfa Detay Verileri"
    sheet2["B1"] = 42

    workbook.save(path)
    return path


def generate_pptx(output_dir: Path) -> Path:
    """Create a realistic PPTX with 2 slides, tables, speaker notes, and Turkish characters."""
    path = output_dir / "sample.pptx"
    presentation = Presentation()

    # Slide 1: Title & Table
    slide1 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide1.shapes.title.text = "RandomOffice Sunum Başlığı (Ş, Ğ, Ç, İ, Ö, Ü)"
    table_shape = slide1.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
    table = table_shape.table
    table.cell(0, 0).text = "Gereksinim"
    table.cell(0, 1).text = "Karşılama"
    table.cell(1, 0).text = "Güvenlik"
    table.cell(1, 1).text = "Sıfır Soket, Defender Dostu"

    # Speaker notes for Slide 1
    notes_slide1 = slide1.notes_slide
    notes_slide1.notes_text_frame.text = "Sunucu için hatırlatma: Bu slayt temel özellikleri özetler."

    # Slide 2: Second slide
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide2.shapes.title.text = "İkinci Slayt: Yol Haritası"
    body_shape = slide2.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "1. Açık ve Ücretsiz Mimari"
    p = tf.add_paragraph()
    p.text = "2. Tam Biçimlendirme Uyumluluğu"

    notes_slide2 = slide2.notes_slide
    notes_slide2.notes_text_frame.text = "Görüşme notu: İkinci slayt planlama aşamalarını anlatır."

    presentation.save(path)
    return path


def main() -> int:
    """Generate all sample files and print their paths."""
    parser = ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("samples"),
        help="directory for generated files (default: samples)",
    )
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    for generator in (generate_docx, generate_xlsx, generate_pptx):
        print(generator(args.output_dir))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

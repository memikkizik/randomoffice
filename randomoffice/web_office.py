"""Interactive Office document serializer and generator for the RandomOffice Web Suite."""

import io
import base64
from pathlib import Path
import re
from typing import Any

from .contracts import OfficeDocument, SupportedFormat
from .errors import DocumentParseError, UnsupportedFormatError
from .editor_formats import (
    apply_cell_style, apply_page_layout, cell_address, read_cell_style,
    read_page_layout, validate_cell_styles, validate_page_layout, workbook_theme_colors,
)
from .ppt_richtext import text_frame_html, write_text_frame
from .word_richtext import append_editor_html, paragraph_html


def open_office_file(source_path: Path) -> dict[str, Any]:
    """Parse an Office document into a rich JSON structure ready for the web editor."""
    suffix = source_path.suffix.lower()

    if suffix == SupportedFormat.DOCX:
        parsed = _parse_docx(source_path)
    elif suffix == SupportedFormat.XLSX:
        parsed = _parse_xlsx(source_path)
    elif suffix == SupportedFormat.PPTX:
        parsed = _parse_pptx(source_path)
    else:
        raise UnsupportedFormatError(f"Unsupported office format: {suffix}")
    canonical = OfficeDocument.from_editor_data(parsed, source_path=source_path)
    return canonical.to_editor_data()


def export_office_file(doc_type: str, data: dict[str, Any] | OfficeDocument) -> bytes:
    """Generate a real DOCX, XLSX, or PPTX file from the web editor data."""
    doc_type = doc_type.lower()
    if isinstance(data, OfficeDocument):
        data = data.to_editor_data()
    _validate_editor_data(doc_type, data)
    if doc_type in ("docx", "word"):
        return _export_docx(data)
    elif doc_type in ("xlsx", "excel", "sheet"):
        return _export_xlsx(data)
    elif doc_type in ("pptx", "powerpoint", "slides"):
        return _export_pptx(data)
    else:
        raise UnsupportedFormatError(f"Cannot export unknown document type: {doc_type}")


def _validate_editor_data(doc_type: str, data: dict[str, Any]) -> None:
    """Reject malformed editor structures before passing them to Office libraries."""
    if not isinstance(data, dict):
        raise ValueError("Editor data must be an object")

    def objects(key: str) -> list[dict[str, Any]]:
        value = data.get(key, [])
        if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
            raise ValueError(f"{key} must be a list of objects")
        return value

    def rows(value: Any) -> None:
        if not isinstance(value, list) or any(not isinstance(row, list) for row in value):
            raise ValueError("Table rows must be lists")
        if sum(len(row) for row in value) > 200_000:
            raise ValueError("Table exceeds 200,000 cells")
        if any(isinstance(cell, (dict, list)) for row in value for cell in row):
            raise ValueError("Table cells must contain scalar values")

    if doc_type in ("word", "docx"):
        if "page_layout" in data:
            validate_page_layout(data["page_layout"])
        for item in objects("items"):
            if item.get("type") == "table":
                rows(item.get("rows", []))
    elif doc_type in ("excel", "xlsx", "sheet"):
        names = set()
        for index, sheet in enumerate(objects("sheets")):
            name = sheet.get("name", f"Sheet{index + 1}")
            if not isinstance(name, str) or not name.strip() or len(name) > 31 or re.search(r"[\\/?*\[\]:]", name):
                raise ValueError("Sheet names must be 1–31 characters without \\ / ? * [ ] :")
            if name.casefold() in names:
                raise ValueError("Sheet names must be unique")
            names.add(name.casefold())
            rows(sheet.get("rows", []))
            if not isinstance(sheet.get("number_formats", {}), dict):
                raise ValueError("number_formats must be an object")
            validate_cell_styles(sheet.get("cell_styles", {}))
            if "freeze_first_row" in sheet and not isinstance(sheet["freeze_first_row"], bool):
                raise ValueError("freeze_first_row must be boolean")
    elif doc_type in ("powerpoint", "pptx", "slides"):
        for slide in objects("slides"):
            for field in ("title_html", "body_html"):
                if field in slide and not isinstance(slide[field], str):
                    raise ValueError(f"Slide {field} must be text")
            paragraphs = slide.get("paragraphs", [])
            if not isinstance(paragraphs, list) or any(not isinstance(text, str) for text in paragraphs):
                raise ValueError("Slide paragraphs must be a list of text")
            tables = slide.get("tables", [])
            if not isinstance(tables, list):
                raise ValueError("Slide tables must be a list")
            for table in tables:
                rows(table)
            images = slide.get("images", [])
            if not isinstance(images, list) or any(not isinstance(image, str) for image in images):
                raise ValueError("Slide images must be a list of data URLs")


# ---------------------------------------------------------------------------
# DOCX Handlers
# ---------------------------------------------------------------------------

def _parse_docx(source_path: Path) -> dict[str, Any]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(str(source_path))
    except Exception as exc:
        raise DocumentParseError(f"Failed to open Word document: {exc}") from exc

    title = doc.core_properties.title if doc.core_properties and doc.core_properties.title else source_path.stem
    items: list[dict[str, Any]] = []

    for child in doc.element.body:
        if child.tag.endswith("p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            style = p.style.name.lower() if p.style else "normal"
            level = "p"
            if "heading 1" in style:
                level = "h1"
            elif "heading 2" in style:
                level = "h2"
            elif "heading 3" in style:
                level = "h3"
            elif "number" in style:
                level = "number"
            elif "bullet" in style or "list" in style:
                level = "bullet"

            items.append({
                "type": "paragraph",
                "text": text,
                "level": level,
                "html": paragraph_html(p),
                "alignment": {0: "left", 1: "center", 2: "right", 3: "justify"}.get(p.alignment, "left"),
            })
        elif child.tag.endswith("tbl"):
            tbl = Table(child, doc)
            grid: list[list[str]] = []
            for row in tbl.rows:
                seen_cells = set()
                row_vals: list[str] = []
                for cell in row.cells:
                    cell_id = id(cell._tc)
                    if cell_id not in seen_cells:
                        seen_cells.add(cell_id)
                        row_vals.append(cell.text.strip())
                if any(row_vals):
                    grid.append(row_vals)
            if grid:
                items.append({
                    "type": "table",
                    "rows": grid,
                })

    return {
        "app": "word",
        "format": ".docx",
        "title": title or "Untitled Document",
        "items": items,
        "page_layout": read_page_layout(doc),
    }


def _export_docx(data: dict[str, Any]) -> bytes:
    from docx import Document

    doc = Document()
    doc.core_properties.title = data.get("title", "RandomOffice Document")
    if "page_layout" in data:
        apply_page_layout(doc, data["page_layout"])

    if "html" in data:
        if not isinstance(data["html"], str):
            raise ValueError("Word editor HTML must be text")
        append_editor_html(doc, data["html"])
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    for item in data.get("items", []):
        itype = item.get("type")
        if itype == "paragraph":
            level = item.get("level", "p")
            text = item.get("text", "")
            if level == "h1":
                doc.add_heading(text, level=1)
            elif level == "h2":
                doc.add_heading(text, level=2)
            elif level == "h3":
                doc.add_heading(text, level=3)
            elif level == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            elif level == "number":
                doc.add_paragraph(text, style="List Number")
            else:
                doc.add_paragraph(text)
            if "html" in item:
                paragraph = doc.paragraphs[-1]
                paragraph.clear()
                from .word_richtext import _write_inline
                from lxml import html
                _write_inline(paragraph, html.fragment_fromstring(item["html"] or "<span></span>", create_parent="span"))
        elif itype == "table":
            rows = item.get("rows", [])
            if rows and len(rows) > 0:
                cols_count = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=cols_count)
                tbl.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, val in enumerate(row):
                        tbl.cell(r_idx, c_idx).text = str(val)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# XLSX Handlers
# ---------------------------------------------------------------------------

def _parse_xlsx(source_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = None
    try:
        try:
            with source_path.open("rb") as stream:
                wb = load_workbook(filename=stream, data_only=False)
        except Exception as exc:
            raise DocumentParseError(f"Failed to open Excel workbook: {exc}") from exc
        sheets: list[dict[str, Any]] = []
        theme = workbook_theme_colors(wb)

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows: list[list[Any]] = []
            max_r = ws.max_row or 1
            max_c = ws.max_column or 1
            if max_r * max_c > 200_000:
                raise DocumentParseError(
                    "Worksheet exceeds the editor limit of 200,000 cells; "
                    "open a smaller workbook to avoid data loss."
                )

            for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=True):
                rows.append([_serialize_cell_value(value) for value in row])

            cell_styles = {}
            for row in ws:
                for cell in row:
                    style = read_cell_style(cell, theme)
                    if style:
                        cell_styles[f"{cell.row - 1}:{cell.column - 1}"] = style
            sheets.append({
                "name": sheetname,
                "rows": rows,
                "number_formats": {f"{cell.row - 1}:{cell.column - 1}": cell.number_format
                                   for row in ws for cell in row if cell.number_format != "General"},
                "cell_styles": cell_styles,
                "freeze_first_row": bool(ws.freeze_panes and ws.freeze_panes != "A1" and
                                         ws[ws.freeze_panes].row > 1),
            })

        title = wb.properties.title if hasattr(wb, "properties") and wb.properties and wb.properties.title else source_path.stem

        return {
            "app": "excel",
            "format": ".xlsx",
            "title": title or "Untitled Workbook",
            "sheets": sheets,
        }
    finally:
        if wb is not None:
            wb.close()


def _export_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    sheets_data = data.get("sheets", [])

    if not sheets_data:
        ws = wb.active
        ws.title = "Sheet1"
    else:
        for idx, s_info in enumerate(sheets_data):
            if idx == 0:
                ws = wb.active
                ws.title = s_info.get("name", "Sheet1")
            else:
                ws = wb.create_sheet(title=s_info.get("name", f"Sheet{idx+1}"))

            for r_idx, row in enumerate(s_info.get("rows", []), start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=_coerce_cell_value(val))
                    number_format = s_info.get("number_formats", {}).get(f"{r_idx - 1}:{c_idx - 1}")
                    if isinstance(number_format, str):
                        ws.cell(row=r_idx, column=c_idx).number_format = number_format
            for key, style in s_info.get("cell_styles", {}).items():
                row, column = cell_address(key)
                apply_cell_style(ws.cell(row=row, column=column), style)
            if s_info.get("freeze_first_row"):
                ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _serialize_cell_value(value: Any) -> Any:
    """Convert an openpyxl cell value into a JSON-safe, lossless value."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _coerce_cell_value(value: Any) -> Any:
    """Preserve formulas and identifiers while restoring obvious numeric input."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text or text.startswith("="):
        return text
    if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", text):
        if len(text.lstrip("-").split(".", 1)[0]) > 1 and text.lstrip("-")[0] == "0":
            return text
        return float(text) if "." in text else int(text)
    return value


# ---------------------------------------------------------------------------
# PPTX Handlers
# ---------------------------------------------------------------------------

def _parse_pptx(source_path: Path) -> dict[str, Any]:
    from pptx import Presentation

    try:
        prs = Presentation(str(source_path))
    except Exception as exc:
        raise DocumentParseError(f"Failed to open PowerPoint presentation: {exc}") from exc

    slides: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        paragraphs: list[str] = []
        tables: list[list[list[str]]] = []
        images: list[str] = []
        body_html: list[str] = []
        title_shape = slide.shapes.title
        title_html = text_frame_html(title_shape.text_frame) if title_shape is not None else ""

        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()

        for shape in slide.shapes:
            if shape.shape_type == 13 and shape.image.content_type in ("image/png", "image/jpeg", "image/gif"):
                images.append(f"data:{shape.image.content_type};base64," + base64.b64encode(shape.image.blob).decode("ascii"))
            elif shape.has_table:
                tbl_rows = []
                for row in shape.table.rows:
                    r_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(r_cells):
                        tbl_rows.append(r_cells)
                if tbl_rows:
                    tables.append(tbl_rows)
            elif shape.has_text_frame:
                if title_shape is not None and shape._element is title_shape._element:
                    continue
                body_html.append(text_frame_html(shape.text_frame))
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        paragraphs.append(t)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slides.append({
            "slide_number": idx,
            "title": slide_title or f"Slide {idx}",
            "paragraphs": paragraphs,
            "tables": tables,
            "notes": notes,
            "images": images,
            "title_html": title_html,
            "body_html": "".join(body_html),
            "theme": _read_slide_theme(slide),
        })

    title = prs.core_properties.title if hasattr(prs, "core_properties") and prs.core_properties.title else source_path.stem

    return {
        "app": "powerpoint",
        "format": ".pptx",
        "title": title or "Untitled Presentation",
        "slides": slides,
    }


def _export_pptx(data: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slides_data = data.get("slides", [])

    if not slides_data:
        # Default blank slide
        prs.slides.add_slide(prs.slide_layouts[0])
    else:
        for s_info in slides_data:
            # Layout 1 is title and content
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)

            from pptx.dml.color import RGBColor
            theme_colors = {"dark": ("1E1E1E", "F3F2F1"), "navy": ("0F2B5C", "FFFFFF"),
                            "emerald": ("0D4D35", "FFFFFF"), "coral": ("5C2014", "FFFFFF")}
            theme = theme_colors.get(s_info.get("theme"))
            if theme:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(theme[0])
            foreground = theme[1] if theme else None

            if slide.shapes.title:
                write_text_frame(slide.shapes.title.text_frame, s_info.get("title_html"),
                                 [s_info.get("title", "")], default_color=foreground)

            # Content placeholder
            content_shape = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if content_shape and content_shape.has_text_frame:
                write_text_frame(content_shape.text_frame, s_info.get("body_html"),
                                 s_info.get("paragraphs", []), default_color=foreground, strip_bullets=True)

            for table_rows in s_info.get("tables", []):
                if not table_rows:
                    continue
                row_count = len(table_rows)
                column_count = max(len(row) for row in table_rows)
                table_shape = slide.shapes.add_table(
                    row_count,
                    column_count,
                    Inches(1),
                    Inches(4.8),
                    Inches(8),
                    Inches(1.5),
                )
                table = table_shape.table
                for row_index, row in enumerate(table_rows):
                    for column_index, value in enumerate(row):
                        table.cell(row_index, column_index).text = str(value)

            # Add speaker notes
            notes = s_info.get("notes")
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            for image in s_info.get("images", []):
                from .word_richtext import decode_image
                slide.shapes.add_picture(io.BytesIO(decode_image(image)), Inches(5), Inches(2), width=Inches(3))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _read_slide_theme(slide) -> str:
    """Recognize only the editor's named solid-background themes."""
    from pptx.enum.dml import MSO_FILL_TYPE

    fill = slide.background.fill
    if fill.type != MSO_FILL_TYPE.SOLID or fill.fore_color.type is None:
        return "default"
    try:
        color = str(fill.fore_color.rgb)
    except AttributeError:
        return "default"
    return {"1E1E1E": "dark", "0F2B5C": "navy", "0D4D35": "emerald", "5C2014": "coral"}.get(color, "default")

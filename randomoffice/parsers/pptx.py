"""PPTX parser backed by python-pptx with table, group shape, and notes extraction."""

from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from ..contracts import ParsedDocument, SupportedFormat
from ..errors import DocumentParseError, ParserDependencyError
from ..performance.guard import FileSecurityGuard
from ._common import modified_at, validate_source


class PptxParser:
    """Extract text from presentation shapes, tables, and notes in slide order."""

    supported_format = SupportedFormat.PPTX

    def __init__(self, security_guard: FileSecurityGuard | None = None) -> None:
        self.security_guard = security_guard or FileSecurityGuard()

    def parse(self, source_path: Path) -> ParsedDocument:
        """Parse PPTX into normalized text and metadata."""
        validate_source(source_path, self.supported_format)
        self.security_guard.validate_office_archive(source_path)

        try:
            from pptx import Presentation
            from pptx.exc import PackageNotFoundError
        except ImportError as exc:
            raise ParserDependencyError(
                "PPTX parsing requires python-pptx; install project dependencies first"
            ) from exc

        try:
            presentation = Presentation(str(source_path))
            parts: list[str] = []

            for slide in presentation.slides:
                self._extract_shapes(slide.shapes, parts)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(f"[Note] {notes}")

            prop_title = (
                presentation.core_properties.title.strip()
                if hasattr(presentation, "core_properties") and presentation.core_properties.title
                else None
            )
            title = prop_title or (parts[0] if parts else None)

            return ParsedDocument(
                source_path=source_path,
                format=self.supported_format,
                title=title,
                text="\n".join(parts),
                modified_at=modified_at(source_path),
            )
        except (
            OSError,
            ValueError,
            KeyError,
            IndexError,
            AttributeError,
            PackageNotFoundError,
            BadZipFile,
        ) as exc:
            raise DocumentParseError(f"Could not parse PPTX: {source_path}") from exc

    def _extract_shapes(self, shapes: Any, parts: list[str]) -> None:
        for shape in shapes:
            if hasattr(shape, "shapes"):
                self._extract_shapes(shape.shapes, parts)
                continue

            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
                continue

            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    parts.append(text)

"""Source-aware Office saves with a byte-exact no-op and conservative overlays.

This module never silently falls back to creating a fresh document. A caller can
offer an explicitly separate Save As operation after a PreservationError.
"""

from copy import deepcopy, copy
from hashlib import sha256
from html import escape
import io
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile, BadZipFile, ZIP_DEFLATED

from lxml import etree, html

from .editor_formats import (apply_cell_style, apply_page_layout, cell_address,
                             validate_cell_styles, validate_page_layout)
from .ppt_richtext import _html_paragraphs, _strip_bullet, _style, write_text_frame
from .word_richtext import append_editor_html, _write_inline


_FORMATS = {"word": ("word", ".docx"), "docx": ("word", ".docx"),
            "excel": ("excel", ".xlsx"), "xlsx": ("excel", ".xlsx"), "sheet": ("excel", ".xlsx"),
            "powerpoint": ("powerpoint", ".pptx"), "pptx": ("powerpoint", ".pptx"), "slides": ("powerpoint", ".pptx")}
_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
       "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
       "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
       "r": "http://schemas.openxmlformats.org/package/2006/relationships"}


class PreservationError(ValueError):
    """The edit cannot safely overwrite the source; explicitly offer Save As."""

    def __init__(self, reason: str, feature: str = "source_preservation"):
        self.reason = reason
        self.requires_save_as = True
        self.loss_report = [{"feature": feature, "message": reason, "severity": "warning"}]
        super().__init__(reason)

    def to_dict(self) -> dict[str, Any]:
        return {"error": self.reason, "requires_save_as": True, "loss_report": self.loss_report}


def _kind(doc_type: str) -> str:
    try:
        return _FORMATS[doc_type.lower().lstrip(".")][0]
    except (KeyError, AttributeError):
        raise PreservationError("Bu dosya türü kaynak korunarak kaydedilemiyor; farklı kaydedin.") from None


def _xml(source: bytes):
    return etree.fromstring(source, parser=etree.XMLParser(resolve_entities=False, no_network=True))


def _parts(source: bytes) -> dict[str, bytes]:
    try:
        with ZipFile(io.BytesIO(source)) as archive:
            return {entry.filename: archive.read(entry) for entry in archive.infolist()}
    except (BadZipFile, KeyError, OSError) as exc:
        raise PreservationError("Kaynak Office arşivi okunamadı; özgün dosyayı yeniden açın.") from exc


def _package(original: bytes, replacements: dict[str, bytes]) -> bytes:
    """Retain untouched archive members and their ZIP metadata."""
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original)) as old, ZipFile(output, "w", compression=ZIP_DEFLATED) as new:
        new.comment = old.comment
        seen = set()
        for info in old.infolist():
            new.writestr(info, replacements.get(info.filename, old.read(info)))
            seen.add(info.filename)
        for name, content in replacements.items():
            if name not in seen:
                new.writestr(name, content)
    return output.getvalue()


def inspect_source(path: Path) -> dict[str, Any]:
    """Describe source features without claiming unsupported editing fidelity."""
    path = Path(path)
    app = _kind(path.suffix)
    source = path.read_bytes()
    parts = _parts(source)
    counts = {"charts": sum(bool(re.fullmatch(r"(?:xl|ppt|word)/charts/chart\d+\.xml", name)) for name in parts),
              "embedded_objects": sum("/embeddings/" in name for name in parts),
              "custom_xml_parts": sum(name.startswith("customXml/") and name.endswith(".xml") for name in parts),
              "digital_signatures": any(name.startswith("_xmlsignatures/") for name in parts)}
    if app == "word":
        root = _xml(parts["word/document.xml"])
        counts.update(paragraphs=len(root.findall("./w:body/w:p", _NS)),
                      tables=len(root.findall("./w:body/w:tbl", _NS)),
                      sections=len(root.findall(".//w:sectPr", _NS)),
                      headers=sum(bool(re.fullmatch(r"word/header\d+\.xml", name)) for name in parts),
                      footers=sum(bool(re.fullmatch(r"word/footer\d+\.xml", name)) for name in parts),
                      fields=len(root.findall(".//w:fldChar", _NS)))
    elif app == "excel":
        sheets = [_xml(content) for name, content in parts.items() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)]
        counts.update(sheets=len(sheets), merged_ranges=sum(len(root.findall(".//s:mergeCell", _NS)) for root in sheets))
    else:
        slides = [_xml(content) for name, content in parts.items() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        counts.update(slides=len(slides), shapes=sum(len(root.findall(".//p:sp", _NS)) for root in slides),
                      pictures=sum(len(root.findall(".//p:pic", _NS)) for root in slides),
                      tables=sum(len(root.findall(".//a:tbl", _NS)) for root in slides))
    loss_report = []
    if counts["digital_signatures"]:
        loss_report.append({"feature": "digital_signature", "severity": "warning",
                            "message": "İmzalı dosyalarda düzenleme dijital imzayı geçersiz kılar; farklı kaydedin."})
    return {"app": app, "format": _FORMATS[app][1], "source_name": path.name,
            "source_fingerprint": sha256(source).hexdigest(), "features": counts,
            "requires_save_as": bool(counts["digital_signatures"]), "loss_report": loss_report}


def _css_state(element, inherited):
    state = _style(element, inherited)
    css = {key.strip().lower(): value.strip() for key, separator, value in
           (part.partition(":") for part in element.get("style", "").split(";")) if separator}
    for field in ("background-color", "text-align"):
        if field in css:
            state[field] = css[field].lower()
    if element.tag == "a":
        state["href"] = element.get("href", "")
    return state


def _normal_style(state):
    result = {}
    for key, value in state.items():
        if key in ("bold", "italic", "underline", "strike") and not value:
            continue
        if key == "text-align" and value == "left":
            continue
        if key == "font_name":
            value = value.lower()
        result[key] = value
    return tuple(sorted(result.items()))


def _inline(element, inherited=None):
    """Canonical text/style tokens, independent of span/b/strong syntax."""
    tag = element.tag.lower() if isinstance(element.tag, str) else ""
    if tag in ("script", "style", "iframe", "object", "svg", "embed"):
        return []
    state = _css_state(element, inherited or {})
    result = []

    def add(text, style):
        if not text:
            return
        signature = _normal_style(style)
        if result and result[-1][0] == "text" and result[-1][2] == signature:
            result[-1] = ("text", result[-1][1]+text, signature)
        else:
            result.append(("text", text, signature))

    if tag == "img":
        return [("image", element.get("src", ""))]
    if tag == "br":
        return [("text", "\n", _normal_style(state))]
    add(element.text, state)
    for child in element:
        for token in _inline(child, state):
            if token[0] == "text":
                add(token[1], dict(token[2]))
            else:
                result.append(token)
        add(child.tail, state)
    return result


def _word_html(data):
    if "html" in data:
        if not isinstance(data["html"], str):
            raise ValueError("Word editor HTML must be text")
        return data["html"]
    blocks = []
    for item in data.get("items", []):
        if item.get("type") == "table":
            blocks.append("<table>"+"".join("<tr>"+"".join("<td>"+escape(str(cell))+"</td>" for cell in row)+"</tr>" for row in item.get("rows", []))+"</table>")
        elif item.get("type") == "paragraph":
            level = item.get("level", "p")
            tag = level if level in ("h1", "h2", "h3") else "li" if level in ("bullet", "number") else "p"
            content = item.get("html") if isinstance(item.get("html"), str) else escape(item.get("text", "")).replace("\n", "<br>")
            alignment = item.get("alignment", "left")
            alignment = alignment if alignment in ("left", "center", "right", "justify") else "left"
            block = f'<{tag} style="text-align:{alignment}">{content}</{tag}>'
            if tag == "li":
                list_tag = "ol" if level == "number" else "ul"
                block = f"<{list_tag}>{block}</{list_tag}>"
            blocks.append(block)
    return "".join(blocks)


def _word_blocks(data):
    parser = html.HTMLParser(no_network=True, recover=True)
    root = html.fragment_fromstring(_word_html(data) or "<p></p>", create_parent="div", parser=parser)
    blocks = []

    def visit(element, list_kind=None):
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in ("script", "style", "iframe", "object"):
            return
        if tag in ("ol", "ul"):
            for child in element:
                visit(child, "number" if tag == "ol" else "bullet")
        elif tag == "table":
            rows = element.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            cells = [[_inline(cell) for cell in row.xpath("./td|./th")] for row in rows]
            blocks.append({"kind": "table", "cells": cells, "element": element})
        elif tag == "div" and any(child.tag in ("p", "div", "table", "ol", "ul") for child in element):
            if element.text and element.text.strip():
                p = html.Element("p"); p.text = element.text; visit(p)
            for child in element:
                visit(child)
        else:
            state = _css_state(element, {})
            alignment = state.get("text-align", "left")
            blocks.append({"kind": "paragraph", "level": list_kind or (tag if tag in ("h1", "h2", "h3") else "p"),
                           "alignment": alignment, "runs": _inline(element), "element": element})
    if root.text and root.text.strip():
        p = html.Element("p"); p.text = root.text; visit(p)
    for child in root:
        visit(child)
    return blocks


def _without_elements(blocks):
    return [{key: value for key, value in block.items() if key != "element"} for block in blocks]


def _coerce(value):
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", text):
            return float(text) if "." in text else int(text)
    return value


def _value(value):
    value = _coerce(value)
    if value is None:
        return None
    return ("bool" if isinstance(value, bool) else "number" if isinstance(value, (float, int)) else "text", value)


def _grid(rows):
    result = {}
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if _value(value) is not None:
                result[f"{row_index}:{column_index}"] = _value(value)
    return result


def _cell_styles(styles):
    return {key: _normal_style(style) for key, style in styles.items() if _normal_style(style)}


def _ppt_text(slide, which):
    field = "title_html" if which == "title" else "body_html"
    if isinstance(slide.get(field), str):
        paragraphs = _html_paragraphs(slide[field])
    else:
        plain = [slide.get("title", "")] if which == "title" else slide.get("paragraphs", slide.get("bullets", []))
        paragraphs = [[(text, {})] for text in plain]
    result = []
    for runs in paragraphs:
        if which == "body":
            runs = _strip_bullet(runs)
        tokens = []
        for text, style in runs:
            signature = _normal_style(style)
            if tokens and tokens[-1][1] == signature:
                tokens[-1] = (tokens[-1][0]+text, signature)
            else:
                tokens.append((text, signature))
        result.append(tokens)
    return result or [[]]


def _semantic(app, data):
    if app == "word":
        return {"blocks": _without_elements(_word_blocks(data)), "page_layout": data.get("page_layout")}
    if app == "excel":
        return [{"name": sheet.get("name", f"Sheet{index+1}"), "rows": _grid(sheet.get("rows", [])),
                 "cell_styles": _cell_styles(sheet.get("cell_styles", {})), "number_formats": sheet.get("number_formats", {}),
                 "freeze_first_row": bool(sheet.get("freeze_first_row", False))}
                for index, sheet in enumerate(data.get("sheets", []))]
    return [{"title": _ppt_text(slide, "title"), "body": _ppt_text(slide, "body"),
             "notes": slide.get("notes", ""), "tables": slide.get("tables", []), "images": slide.get("images", []),
             "theme": slide.get("theme", "default")}
            for slide in data.get("slides", [])]


def semantically_unchanged(doc_type: str, payload: dict, baseline_editor_data: dict) -> bool:
    """Compare editable content while ignoring UI filenames/envelope metadata."""
    return _semantic(_kind(doc_type), payload) == _semantic(_kind(doc_type), baseline_editor_data)


def export_preserving_source(doc_type: str, payload: dict, source_bytes: bytes,
                             baseline_editor_data: dict) -> bytes:
    app = _kind(doc_type)
    if not isinstance(payload, dict) or not isinstance(baseline_editor_data, dict) or not isinstance(source_bytes, bytes):
        raise PreservationError("Kaynak kaydetme verisi geçersiz; dosyayı yeniden açın.")
    if semantically_unchanged(app, payload, baseline_editor_data):
        return source_bytes
    parts = _parts(source_bytes)
    if any(name.startswith("_xmlsignatures/") for name in parts):
        raise PreservationError("Düzenleme dijital imzayı geçersiz kılar. Özgün dosyayı korumak için farklı kaydedin.", "digital_signature")
    try:
        if app == "excel":
            return _overlay_xlsx(payload, source_bytes, baseline_editor_data, parts)
        if app == "word":
            return _overlay_word(payload, source_bytes, baseline_editor_data)
        return _overlay_ppt(payload, source_bytes, baseline_editor_data)
    except PreservationError:
        raise
    except Exception as exc:
        raise PreservationError("Bu düzenleme kaynak yapısı korunarak kaydedilemedi. Özgün dosyanızı koruyarak farklı kaydedin.") from exc


def _serialize(root):
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def _sheet_paths(parts):
    import posixpath
    relationships = _xml(parts["xl/_rels/workbook.xml.rels"])
    targets = {rel.get("Id"): rel.get("Target") for rel in relationships}
    root = _xml(parts["xl/workbook.xml"])
    result = []
    for sheet in root.findall("s:sheets/s:sheet", _NS):
        target = targets[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
        result.append(target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target)))
    return result


def _overlay_xlsx(payload, source, baseline, original_parts):
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import PatternFill

    sheets, old_sheets = payload.get("sheets", []), baseline.get("sheets", [])
    if len(sheets) != len(old_sheets) or [s.get("name") for s in sheets] != [s.get("name") for s in old_sheets]:
        raise PreservationError("Sayfa ekleme, silme veya ad değiştirme kaynak başvurularını etkiler. Bu değişikliği farklı kaydedin.", "worksheet_structure")
    workbook = load_workbook(io.BytesIO(source), data_only=False, keep_links=True)
    changed = {}
    freeze_changes = set()
    any_styles = False
    try:
        for index, (new, old) in enumerate(zip(sheets, old_sheets)):
            sheet = workbook.worksheets[index]
            if sheet.title != old.get("name"):
                raise PreservationError("Kaynak sayfa sırası değişmiş; dosyayı yeniden açın.", "worksheet_structure")
            validate_cell_styles(new.get("cell_styles", {}))
            new_grid, old_grid = _grid(new.get("rows", [])), _grid(old.get("rows", []))
            cells = set(new_grid) | set(old_grid)
            changed[index] = set()
            for key in cells:
                if new_grid.get(key) == old_grid.get(key):
                    continue
                row, column = cell_address(key)
                cell = sheet.cell(row=row, column=column)
                if isinstance(cell, MergedCell):
                    raise PreservationError("Birleştirilmiş hücrelerin iç bölümleri değiştirilemiyor; farklı kaydedin.", "merged_cells")
                cell.value = new_grid[key][1] if key in new_grid else None
                changed[index].add(cell.coordinate)
            new_styles, old_styles = new.get("cell_styles", {}), old.get("cell_styles", {})
            for key in set(new_styles) | set(old_styles):
                style, old_style = new_styles.get(key, {}), old_styles.get(key, {})
                if _normal_style(style) == _normal_style(old_style):
                    continue
                row, column = cell_address(key)
                cell = sheet.cell(row=row, column=column)
                if isinstance(cell, MergedCell):
                    raise PreservationError("Birleştirilmiş hücre biçimi yalnızca sol üst hücreden değiştirilebilir.", "merged_cells")
                removed = set(old_style)-set(style)
                font = copy(cell.font)
                defaults = workbook._fonts[0]
                for field, attribute in (("bold", "bold"), ("italic", "italic"), ("underline", "underline"),
                                         ("strike", "strike"), ("font_name", "name"), ("font_size", "size"), ("color", "color")):
                    if field in removed:
                        setattr(font, attribute, copy(getattr(defaults, attribute)))
                cell.font = font
                if "background" in removed:
                    cell.fill = PatternFill()
                if "alignment" in removed:
                    alignment = copy(cell.alignment); alignment.horizontal = None; cell.alignment = alignment
                apply_cell_style(cell, style)
                changed[index].add(cell.coordinate)
                any_styles = True
            formats, old_formats = new.get("number_formats", {}), old.get("number_formats", {})
            for key in set(formats) | set(old_formats):
                if formats.get(key, "General") == old_formats.get(key, "General"):
                    continue
                row, column = cell_address(key)
                cell = sheet.cell(row=row, column=column)
                cell.number_format = formats.get(key, "General")
                changed[index].add(cell.coordinate); any_styles = True
            if bool(new.get("freeze_first_row")) != bool(old.get("freeze_first_row")):
                sheet.freeze_panes = "A2" if new.get("freeze_first_row") else None
                freeze_changes.add(index)
        output = io.BytesIO(); workbook.save(output)
        generated = _parts(output.getvalue())
        paths = _sheet_paths(original_parts)
        generated_paths = _sheet_paths(generated)
        replacements = {}
        for index, path in enumerate(paths):
            if not changed[index] and index not in freeze_changes:
                continue
            original = _xml(original_parts[path]); updated = _xml(generated[generated_paths[index]])
            old_data = original.find("s:sheetData", _NS)
            new_data = updated.find("s:sheetData", _NS)
            old_rows = {row.get("r"): row for row in old_data}
            new_cells = {cell.get("r"): cell for row in new_data for cell in row}
            for address in changed[index]:
                row_number = re.search(r"\d+$", address)[0]
                row = old_rows.get(row_number)
                if row is None:
                    row = etree.Element(f"{{{_NS['s']}}}row", r=row_number)
                    before = next((item for item in old_data if int(item.get("r")) > int(row_number)), None)
                    old_data.insert(list(old_data).index(before) if before is not None else len(old_data), row)
                    old_rows[row_number] = row
                old_cell = next((cell for cell in row if cell.get("r") == address), None)
                replacement = new_cells.get(address)
                if old_cell is not None:
                    location = list(row).index(old_cell)
                    # Unsupported per-cell extension metadata cannot be rebuilt.
                    if any(etree.QName(child).localname not in ("f", "v", "is") for child in old_cell):
                        raise PreservationError("Bu hücre gelişmiş Excel verisi içeriyor; değişikliği farklı kaydedin.", "cell_extensions")
                    row.remove(old_cell)
                    if replacement is not None:
                        row.insert(location, deepcopy(replacement))
                elif replacement is not None:
                    row.append(deepcopy(replacement))
                    def column(cell):
                        letters = re.match(r"[A-Z]+", cell.get("r"))[0]
                        return sum((ord(char)-64)*26**index for index, char in enumerate(reversed(letters)))
                    row[:] = sorted(row, key=column)
            if index in freeze_changes:
                old_views = original.find("s:sheetViews", _NS)
                new_views = updated.find("s:sheetViews", _NS)
                if old_views is not None and new_views is not None:
                    original.replace(old_views, deepcopy(new_views))
            dimension = original.find("s:dimension", _NS)
            new_dimension = updated.find("s:dimension", _NS)
            if dimension is not None and new_dimension is not None:
                dimension.set("ref", new_dimension.get("ref"))
            replacements[path] = _serialize(original)
        if any_styles:
            # Existing style indexes remain stable in openpyxl; only new styles
            # are appended. Preserve unknown style extensions from the source.
            styles = _xml(generated["xl/styles.xml"])
            old_styles_xml = _xml(original_parts["xl/styles.xml"])
            for child in old_styles_xml:
                if not any(etree.QName(candidate).localname == etree.QName(child).localname for candidate in styles):
                    styles.append(deepcopy(child))
            replacements["xl/styles.xml"] = _serialize(styles)
        return _package(source, replacements)
    finally:
        workbook.close()


def _union_saved(source, saved):
    """OPC libraries retain related parts; retain unreferenced custom parts too."""
    return _package(source, _parts(saved))


def _simple_word_paragraph(paragraph):
    allowed = {"pPr", "r", "rPr", "t", "br", "tab", "hyperlink"}
    # pPr may contain section, numbering and paragraph layout objects; keep it.
    for child in paragraph._p:
        if etree.QName(child).localname == "pPr":
            continue
        for node in child.iter():
            localname = etree.QName(node).localname
            if localname in allowed:
                continue
            if any(etree.QName(ancestor).localname == "rPr" for ancestor in node.iterancestors()):
                continue
            return False
    return True


def _overlay_word(payload, source, baseline):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    old_blocks, new_blocks = _word_blocks(baseline), _word_blocks(payload)
    if len(old_blocks) != len(new_blocks) or [b["kind"] for b in old_blocks] != [b["kind"] for b in new_blocks]:
        raise PreservationError("Paragraf veya tablo yapısı değişti. Üstbilgi ve özel nesneleri korumak için farklı kaydedin.", "word_structure")
    document = Document(io.BytesIO(source))
    source_blocks = []
    for element in document.element.body:
        tag = etree.QName(element).localname
        if tag == "p":
            source_blocks.append(Paragraph(element, document))
        elif tag == "tbl":
            table = Table(element, document)
            if any(cell.text.strip() for row in table.rows for cell in row.cells):
                source_blocks.append(table)
    if len(source_blocks) != len(old_blocks):
        raise PreservationError("Kaynak Word blokları güvenle eşleştirilemiyor; farklı kaydedin.", "word_structure")
    for current, old, new in zip(source_blocks, old_blocks, new_blocks):
        if _without_elements([old]) == _without_elements([new]):
            continue
        if new["kind"] == "paragraph":
            if not _simple_word_paragraph(current):
                raise PreservationError("Düzenlenen paragraf alan, yer imi, resim veya özel nesne içeriyor; farklı kaydedin.", "word_inline_objects")
            for child in list(current._p):
                if etree.QName(child).localname != "pPr":
                    current._p.remove(child)
            _write_inline(current, deepcopy(new["element"]))
            if new["level"] != old["level"]:
                names = {"p": "Normal", "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "bullet": "List Bullet", "number": "List Number"}
                current.style = names[new["level"]]
            if new["alignment"] != old["alignment"]:
                current.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                     "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[new["alignment"]]
        else:
            if current._tbl.xpath(".//w:gridSpan | .//w:vMerge") or len(current.rows) != len(new["cells"]) or any(len(row.cells) != len(cells) for row, cells in zip(current.rows, new["cells"])):
                raise PreservationError("Birleştirilmiş ya da boyutu değişen Word tablosu güvenle eşleştirilemiyor; farklı kaydedin.", "word_tables")
            html_rows = new["element"].xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr")
            for r, row in enumerate(current.rows):
                for c, cell in enumerate(row.cells):
                    if new["cells"][r][c] == old["cells"][r][c]:
                        continue
                    if len(cell.paragraphs) != 1 or not _simple_word_paragraph(cell.paragraphs[0]):
                        raise PreservationError("Düzenlenen tablo hücresi özel yapı içeriyor; farklı kaydedin.", "word_tables")
                    paragraph = cell.paragraphs[0]
                    for child in list(paragraph._p):
                        if etree.QName(child).localname != "pPr":
                            paragraph._p.remove(child)
                    _write_inline(paragraph, deepcopy(html_rows[r].xpath("./td|./th")[c]))
    if payload.get("page_layout") != baseline.get("page_layout") and "page_layout" in payload:
        validate_page_layout(payload["page_layout"]); apply_page_layout(document, payload["page_layout"])
    output = io.BytesIO(); document.save(output)
    return _union_saved(source, output.getvalue())


def _overlay_ppt(payload, source, baseline):
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    slides, old_slides = payload.get("slides", []), baseline.get("slides", [])
    presentation = Presentation(io.BytesIO(source))
    if len(slides) != len(old_slides) or len(slides) != len(presentation.slides):
        raise PreservationError("Slayt ekleme veya silme kaynak nesnelerini etkiler; sunumu farklı kaydedin.", "slide_structure")
    themes = {"dark": ("1E1E1E", "F3F2F1"), "navy": ("0F2B5C", "FFFFFF"),
              "emerald": ("0D4D35", "FFFFFF"), "coral": ("5C2014", "FFFFFF"), "default": ("FFFFFF", "201F1E")}
    for slide, new, old in zip(presentation.slides, slides, old_slides):
        if new.get("images", []) != old.get("images", []):
            raise PreservationError("Resim ekleme, silme veya sıralama kaynak yerleşimini etkiler; farklı kaydedin.", "slide_media")
        title = slide.shapes.title
        if _ppt_text(new, "title") != _ppt_text(old, "title"):
            if title is None:
                raise PreservationError("Bu slaytın düzenlenebilir başlık alanı yok; farklı kaydedin.", "slide_title")
            write_text_frame(title.text_frame, new.get("title_html"), [new.get("title", "")])
        if _ppt_text(new, "body") != _ppt_text(old, "body"):
            frames = [shape.text_frame for shape in slide.shapes if shape.has_text_frame and
                      (title is None or shape._element is not title._element)]
            if len(frames) != 1:
                raise PreservationError("Bu slaytta birden çok metin kutusu var; konumları korumak için farklı kaydedin.", "slide_text_boxes")
            write_text_frame(frames[0], new.get("body_html"), new.get("paragraphs", new.get("bullets", [])), strip_bullets=True)
        if new.get("tables", []) != old.get("tables", []):
            tables = [shape.table for shape in slide.shapes if shape.has_table]
            if len(tables) != len(new.get("tables", [])):
                raise PreservationError("Slayt tablo yapısı değişti; farklı kaydedin.", "slide_tables")
            for table, rows in zip(tables, new.get("tables", [])):
                if len(table.rows) != len(rows) or any(len(row) != len(table.columns) for row in rows):
                    raise PreservationError("Slayt tablo boyutları değişti; farklı kaydedin.", "slide_tables")
                for r, values in enumerate(rows):
                    for c, value in enumerate(values):
                        if table.cell(r, c).is_merge_origin or table.cell(r, c).is_spanned:
                            raise PreservationError("Birleştirilmiş slayt tablosu düzenlemesini farklı kaydedin.", "slide_tables")
                        if table.cell(r, c).text != str(value):
                            table.cell(r, c).text = str(value)
        if new.get("notes", "") != old.get("notes", ""):
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is None:
                raise PreservationError("Bu slaydın not alanı güvenle değiştirilemiyor; farklı kaydedin.", "slide_notes")
            notes_frame.text = new.get("notes", "")
        if new.get("theme", "default") != old.get("theme", "default"):
            theme = themes.get(new.get("theme", "default"))
            if theme is None:
                raise PreservationError("Bu sunum teması kaynak üzerine uygulanamıyor; farklı kaydedin.", "slide_theme")
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(theme[0])
    output = io.BytesIO(); presentation.save(output)
    return _union_saved(source, output.getvalue())

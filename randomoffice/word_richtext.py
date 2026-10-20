"""Small, explicit HTML/DOCX bridge for the editor's basic formatting."""

import base64
import binascii
import io
import re
from html import escape

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from lxml import html


def decode_image(source: str) -> bytes:
    match = re.fullmatch(r"data:image/(?:png|jpeg|gif);base64,([A-Za-z0-9+/=]+)", source)
    if not match or len(match[1]) > 14_000_000:
        raise ValueError("Invalid embedded image or image exceeds 10 MB")
    try:
        return base64.b64decode(match[1], validate=True)
    except binascii.Error as exc:
        raise ValueError("Invalid embedded image") from exc


def paragraph_html(paragraph) -> str:
    """Emit only known formatting and embedded images, never document HTML."""
    parts = []
    from docx.text.run import Run
    runs = []
    for child in paragraph._p:
        if child.tag == qn("w:r"):
            runs.append((Run(child, paragraph), None))
        elif child.tag == qn("w:hyperlink"):
            relationship = paragraph.part.rels.get(child.get(qn("r:id")))
            target = relationship.target_ref if relationship is not None else None
            for run in child.findall(qn("w:r")):
                runs.append((Run(run, paragraph), target))
    for run, target in runs:
        text = escape(run.text).replace("\n", "<br>")
        styles = []
        if run.bold:
            styles.append("font-weight:bold")
        if run.italic:
            styles.append("font-style:italic")
        decorations = []
        if run.underline:
            decorations.append("underline")
        if run.font.strike:
            decorations.append("line-through")
        if decorations:
            styles.append("text-decoration:" + " ".join(decorations))
        if run.font.color and run.font.color.rgb:
            styles.append(f"color:#{run.font.color.rgb}")
        if run.font.size:
            styles.append(f"font-size:{run.font.size.pt}pt")
        if run.font.name:
            styles.append("font-family:" + escape(run.font.name, quote=True))
        if styles:
            text = f'<span style="{";".join(styles)}">{text}</span>'
        if target and re.match(r"^(https?://|mailto:)", target, re.I):
            text = f'<a href="{escape(target, quote=True)}">{text}</a>'
        parts.append(text)
        for blip in run._r.iter(qn("a:blip")):
            relationship = blip.get(qn("r:embed"))
            if relationship in paragraph.part.related_parts:
                image = paragraph.part.related_parts[relationship]
                if image.content_type in ("image/png", "image/jpeg", "image/gif"):
                    encoded = base64.b64encode(image.blob).decode("ascii")
                    parts.append(f'<img src="data:{image.content_type};base64,{encoded}" alt="">')
    return "".join(parts)


def _style_map(value):
    return {key.strip().lower(): val.strip() for key, _, val in
            (part.partition(":") for part in value.split(";")) if key and val}


def _color(value):
    if re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        return RGBColor.from_string(value[1:])
    match = re.fullmatch(r"rgb\(\s*(\d+),\s*(\d+),\s*(\d+)\s*\)", value)
    if match and all(int(component) <= 255 for component in match.groups()):
        return RGBColor(*(int(component) for component in match.groups()))
    return None


def _write_inline(paragraph, element, inherited=None):
    state = dict(inherited or {})
    tag = element.tag.lower() if isinstance(element.tag, str) else ""
    if tag in ("script", "style", "iframe", "object"):
        return
    css = _style_map(element.get("style", ""))
    state.update(css)
    if tag in ("b", "strong"):
        state["font-weight"] = "bold"
    if tag in ("i", "em"):
        state["font-style"] = "italic"
    if tag == "u":
        state["text-decoration"] = "underline"
    if tag in ("s", "strike", "del"):
        state["text-decoration"] = "line-through"
    if tag == "font":
        state["color"] = element.get("color", state.get("color", ""))
        state["font-family"] = element.get("face", state.get("font-family", ""))
        sizes = {"1": "8pt", "2": "10pt", "3": "12pt", "4": "14pt", "5": "18pt", "6": "24pt", "7": "36pt"}
        if element.get("size") in sizes:
            state["font-size"] = sizes[element.get("size")]
    if tag == "br":
        paragraph.add_run().add_break()
    if tag == "a" and re.match(r"^(https?://|mailto:)", element.get("href", ""), re.I):
        count = len(paragraph.runs)
        # Build the link text using the same formatting path, then move the
        # resulting runs under a Word hyperlink relationship.
        wrapper = html.Element("span")
        wrapper.text = element.text
        for child in element:
            wrapper.append(child)
        _write_inline(paragraph, wrapper, state)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(element.get("href"), RT.HYPERLINK, is_external=True))
        for run in paragraph.runs[count:]:
            link.append(run._r)
        paragraph._p.append(link)
        return
    if tag == "img":
        source = element.get("src", "")
        if source.startswith("data:image/"):
            try:
                blob = decode_image(source)
                paragraph.add_run().add_picture(io.BytesIO(blob), width=Inches(5.5))
            except (ValueError, binascii.Error, UnrecognizedImageError) as exc:
                raise ValueError("Invalid embedded image") from exc
        return

    def add_text(text):
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = state.get("font-weight", "") in ("bold", "700", "800", "900")
        run.italic = state.get("font-style") == "italic"
        decoration = state.get("text-decoration", "") + state.get("text-decoration-line", "")
        run.underline = "underline" in decoration
        run.font.strike = "line-through" in decoration
        color = _color(state.get("color", ""))
        if color:
            run.font.color.rgb = color
        if state.get("font-family"):
            run.font.name = state["font-family"].split(",")[0].strip(" '\"")
        size = re.fullmatch(r"([\d.]+)(pt|px)", state.get("font-size", ""))
        if size:
            points = float(size[1]) * (0.75 if size[2] == "px" else 1)
            run.font.size = Pt(max(1, min(points, 400)))

    add_text(element.text)
    for child in element:
        _write_inline(paragraph, child, state)
        add_text(child.tail)


def append_editor_html(document, source: str) -> None:
    """Read the editor's local HTML. Never resolve external URLs or resources."""
    if len(source) > 20_000_000:
        raise ValueError("Word editor content exceeds 20 MB")
    root = html.fragment_fromstring(source or "<p></p>", create_parent="div")
    if root.text:
        document.add_paragraph(root.text)

    def add_block(element, list_style=None):
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in ("script", "style", "iframe", "object"):
            return
        if tag in ("ul", "ol"):
            for child in element:
                add_block(child, "List Bullet" if tag == "ul" else "List Number")
            return
        if tag == "table":
            rows = element.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            column_count = max((len(row.xpath("./td | ./th")) for row in rows), default=0)
            if rows and column_count:
                table = document.add_table(rows=len(rows), cols=column_count)
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for column_index, cell in enumerate(row.xpath("./td | ./th")):
                        _write_inline(table.cell(row_index, column_index).paragraphs[0], cell)
            return
        if tag == "div" and any(child.tag in ("p", "div", "table", "ul", "ol") for child in element):
            if element.text:
                document.add_paragraph(element.text)
            for child in element:
                add_block(child)
                if child.tail:
                    document.add_paragraph(child.tail)
            return
        paragraph = document.add_heading(level=int(tag[1])) if tag in ("h1", "h2", "h3") else document.add_paragraph(style=list_style)
        alignment = _style_map(element.get("style", "")).get("text-align", "")
        if alignment in ("left", "center", "right", "justify"):
            paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                   "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[alignment]
        _write_inline(paragraph, element)

    for element in root:
        add_block(element)
        if element.tail and element.tail.strip():
            document.add_paragraph(element.tail)

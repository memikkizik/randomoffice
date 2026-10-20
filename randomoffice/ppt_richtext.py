"""Safe local HTML/PowerPoint text bridge for basic editor formatting."""

from html import escape
import re
from typing import Any

from lxml import html
from pptx.dml.color import RGBColor
from pptx.util import Pt


_BLOCKS = {"p", "div", "h1", "h2", "h3", "h4", "li", "blockquote"}
_ALLOWED = _BLOCKS | {"span", "b", "strong", "i", "em", "u", "s", "strike", "del", "br", "font", "a", "ul", "ol"}
_SKIP = {"script", "style", "iframe", "object", "embed", "svg", "math", "img", "link"}


def _color(value: str) -> str | None:
    if re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        return value[1:].upper()
    if re.fullmatch(r"#[0-9a-fA-F]{3}", value):
        return "".join(component*2 for component in value[1:]).upper()
    match = re.fullmatch(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", value)
    if match and all(int(component) <= 255 for component in match.groups()):
        return "".join(f"{int(component):02X}" for component in match.groups())
    return None


def _style(element, inherited: dict[str, Any]) -> dict[str, Any]:
    state = dict(inherited)
    tag = element.tag.lower() if isinstance(element.tag, str) else ""
    css = {key.strip().lower(): value.strip() for key, separator, value in
           (part.partition(":") for part in element.get("style", "").split(";")) if separator}
    if tag in ("b", "strong"):
        state["bold"] = True
    if tag in ("i", "em"):
        state["italic"] = True
    if tag == "u":
        state["underline"] = True
    if tag in ("s", "strike", "del"):
        state["strike"] = True
    if "font-weight" in css:
        state["bold"] = css["font-weight"] in ("bold", "bolder", "600", "700", "800", "900")
    if "font-style" in css:
        state["italic"] = css["font-style"] in ("italic", "oblique")
    decoration = css.get("text-decoration", "") + " " + css.get("text-decoration-line", "")
    if decoration.strip():
        state["underline"] = "underline" in decoration
        state["strike"] = "line-through" in decoration
    color = _color(css.get("color", element.get("color", "") if tag == "font" else ""))
    if color:
        state["color"] = color
    name = css.get("font-family", element.get("face", "") if tag == "font" else "")
    if name:
        state["font_name"] = name.split(",")[0].strip(" '\"")[:200]
    size = css.get("font-size", "")
    if tag == "font" and not size:
        size = {"1": "8pt", "2": "10pt", "3": "12pt", "4": "14pt", "5": "18pt", "6": "24pt", "7": "36pt"}.get(element.get("size"), "")
    match = re.fullmatch(r"(\d+(?:\.\d+)?)(pt|px)", size)
    if match:
        state["font_size"] = max(1, min(400, float(match[1])*(0.75 if match[2] == "px" else 1)))
    return state


def _html_paragraphs(source: str) -> list[list[tuple[str, dict[str, Any]]]]:
    if len(source) > 20_000_000:
        raise ValueError("PowerPoint editor HTML exceeds 20 MB")
    parser = html.HTMLParser(no_network=True, recover=True)
    root = html.fragment_fromstring(source or "<p></p>", create_parent="div", parser=parser)
    paragraphs: list[list[tuple[str, dict[str, Any]]]] = []
    current: list[tuple[str, dict[str, Any]]] = []

    def flush(force=False):
        nonlocal current
        if current or force:
            paragraphs.append(current)
            current = []

    def text(value, state):
        if value:
            current.append((value, state))

    def visit(element, inherited):
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in _SKIP or tag not in _ALLOWED:
            return
        state = _style(element, inherited)
        block = tag in _BLOCKS
        if block:
            flush()
        before = len(paragraphs)
        if tag == "br":
            text("\v", state)
        else:
            text(element.text, state)
            for child in element:
                visit(child, state)
                text(child.tail, state)
        if block:
            flush(force=len(paragraphs) == before)

    if root.text and root.text.strip():
        text(root.text, {})
    for element in root:
        visit(element, {})
        # Pretty-print whitespace between blocks is not another paragraph.
        if element.tail and element.tail.strip():
            text(element.tail, {})
    flush()
    return paragraphs or [[]]


def _strip_bullet(runs: list[tuple[str, dict[str, Any]]]) -> list[tuple[str, dict[str, Any]]]:
    match = re.match(r"^[ \t]*•[ \t]?", "".join(text for text, _ in runs))
    remaining = len(match[0]) if match else 0
    result = []
    for text, state in runs:
        trim = min(len(text), remaining)
        remaining -= trim
        if text[trim:]:
            result.append((text[trim:], state))
    return result


def write_text_frame(frame, source: str | None, plain: list[str], *, default_color: str | None = None,
                     strip_bullets: bool = False) -> None:
    """Write known text attributes; HTML URLs/images are never resolved."""
    paragraphs = _html_paragraphs(source) if source is not None else [[(text, {})] for text in plain]
    frame.clear()
    for index, runs in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if strip_bullets:
            runs = _strip_bullet(runs)
        for text, state in runs:
            # Soft breaks are supported by OOXML and remain within one bullet.
            for part_index, part in enumerate(text.split("\v")):
                if part_index:
                    paragraph.add_line_break()
                if not part:
                    continue
                run = paragraph.add_run()
                run.text = part
                for field in ("bold", "italic", "underline"):
                    if field in state:
                        setattr(run.font, field, state[field])
                if "strike" in state:
                    run.font._rPr.set("strike", "sngStrike" if state["strike"] else "noStrike")
                color = state.get("color", default_color)
                if color:
                    run.font.color.rgb = RGBColor.from_string(color)
                if "font_size" in state:
                    run.font.size = Pt(state["font_size"])
                if state.get("font_name"):
                    run.font.name = state["font_name"]


def text_frame_html(frame) -> str:
    """Generate a restricted HTML fragment from PowerPoint runs."""
    result = []
    for paragraph in frame.paragraphs:
        parts = []
        runs = iter(paragraph.runs)
        for child in paragraph._p:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "br":
                parts.append("<br>")
            elif tag == "fld":
                # Dynamic fields are represented as their safe displayed text.
                field_text = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t")
                if field_text is not None:
                    parts.append(escape(field_text.text or ""))
            elif tag == "r":
                run = next(runs)
                styles = []
                for field, name, enabled, disabled in (("bold", "font-weight", "bold", "normal"),
                                                       ("italic", "font-style", "italic", "normal")):
                    value = getattr(run.font, field)
                    if value is not None:
                        styles.append(f"{name}:{enabled if value else disabled}")
                if run.font.underline:
                    styles.append("text-decoration:underline")
                if run.font._rPr.get("strike") in ("sngStrike", "dblStrike"):
                    styles.append("text-decoration-line:line-through" + (" underline" if run.font.underline else ""))
                color = run.font.color
                if color.type is not None:
                    try:
                        rgb = color.rgb
                    except AttributeError:
                        rgb = None
                    if rgb:
                        styles.append(f"color:#{rgb}")
                if run.font.size:
                    styles.append(f"font-size:{run.font.size.pt:g}pt")
                if run.font.name:
                    name = re.sub(r"[;:{}<>\"'\\\r\n]", "", run.font.name)
                    styles.append(f"font-family:{name}")
                text = escape(run.text).replace("\n", "<br>").replace("\v", "<br>")
                parts.append(f'<span style="{escape(";".join(styles), quote=True)}">{text}</span>' if styles else text)
        result.append("<p>" + "".join(parts) + "</p>")
    return "".join(result)

"""Tests for scripts/generate_samples.py verifying realistic Office sample generation."""

from pathlib import Path
import sys

# Ensure project root is in sys.path so scripts can be imported
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from scripts.generate_samples import generate_docx, generate_xlsx, generate_pptx


def test_generate_docx_realistic(tmp_path: Path):
    path = generate_docx(tmp_path)
    assert path.exists()

    doc = Document(path)
    assert doc.paragraphs[0].text == "RandomOffice Örnek Türkçe Belge"

    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "Şemsiye, Çağlayan, İğne, Ömür, Çözüm" in full_text
    assert "Numaralandırılmış Madde 1" in full_text

    # Bold and italic runs
    all_runs = [r for p in doc.paragraphs for r in p.runs]
    assert any(r.bold and "Kalın metin" in r.text for r in all_runs)
    assert any(r.italic and "İtalik metin" in r.text for r in all_runs)

    # Table
    assert len(doc.tables) >= 1
    assert doc.tables[0].cell(0, 0).text == "Alan / Başlık"
    assert "Kusursuz" in doc.tables[0].cell(1, 1).text


def test_generate_xlsx_realistic(tmp_path: Path):
    path = generate_xlsx(tmp_path)
    assert path.exists()

    wb = load_workbook(path)
    assert "Özet" in wb.sheetnames
    assert "Detaylar" in wb.sheetnames

    sheet1 = wb["Özet"]
    assert sheet1["A3"].value == "00123"
    assert sheet1["C4"].value == 0
    assert sheet1["C6"].value == "=SUM(C3:C5)"
    assert sheet1["AA1"].value == "Geniş Sütun Verisi"
    assert "₺" in (sheet1["D3"].number_format or "")

    sheet2 = wb["Detaylar"]
    assert sheet2["B1"].value == 42
    wb.close()


def test_generate_pptx_realistic(tmp_path: Path):
    path = generate_pptx(tmp_path)
    assert path.exists()

    prs = Presentation(path)
    assert len(prs.slides) == 2

    slide1 = prs.slides[0]
    assert "Ş, Ğ, Ç, İ, Ö, Ü" in slide1.shapes.title.text
    assert slide1.has_notes_slide
    assert "hatırlatma" in slide1.notes_slide.notes_text_frame.text
    assert any(shape.has_table for shape in slide1.shapes)

    slide2 = prs.slides[1]
    assert "Yol Haritası" in slide2.shapes.title.text
    assert slide2.has_notes_slide
    assert "planlama aşamalarını" in slide2.notes_slide.notes_text_frame.text


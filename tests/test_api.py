import io
from docx import Document
from fastapi.testclient import TestClient
import pytest

from randomoffice.api.app import app


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_health_endpoint(client: TestClient) -> None:
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "uptime_seconds" in data
    assert ".docx" in data["supported_formats"]
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "SAMEORIGIN"


def test_web_ui_is_self_contained(client: TestClient) -> None:
    response = client.get("/")

    assert response.status_code == 200
    assert "fonts.googleapis.com" not in response.text
    assert "fonts.gstatic.com" not in response.text


def test_manifest_is_available(client: TestClient) -> None:
    response = client.get("/manifest.json")

    assert response.status_code == 200
    assert response.json()["short_name"] == "RandomOffice"
    assert response.json()["icons"][0]["src"] == "/static/randomoffice-logo.svg"


def test_brand_logo_is_served_locally(client: TestClient) -> None:
    response = client.get("/static/randomoffice-logo.svg")

    assert response.status_code == 200
    assert "RandomOffice" in response.text


def test_api_parse_docx(client: TestClient) -> None:
    doc = Document()
    doc.core_properties.title = "API Test Title"
    doc.add_paragraph("Testing API extraction.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    response = client.post(
        "/api/v1/parse",
        files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["title"] == "API Test Title"
    assert "Testing API extraction." in data["text"]


def test_api_convert_docx_to_markdown(client: TestClient) -> None:
    doc = Document()
    doc.add_heading("API Heading", level=1)
    doc.add_paragraph("Paragraph in API.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    response = client.post(
        "/api/v1/convert",
        files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"target": "markdown"},
    )
    assert response.status_code == 200
    data = response.json()
    assert "# API Heading" in data["content"]
    assert "Paragraph in API." in data["content"]


def test_api_rejects_unsupported_file(client: TestClient) -> None:
    response = client.post(
        "/api/v1/parse",
        files={"file": ("test.pdf", b"%PDF-1.4...", "application/pdf")},
    )
    assert response.status_code == 400
    assert "Unsupported file extension" in response.json()["detail"]


def test_api_enforces_upload_limit(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RANDOMOFFICE_MAX_UPLOAD_BYTES", "16")
    response = client.post(
        "/api/v1/parse",
        files={"file": ("test.docx", b"x" * 17, "application/octet-stream")},
    )

    assert response.status_code == 400
    assert "Upload exceeds limit" in response.json()["detail"]


def test_docs_and_schema_are_hidden_by_default(client: TestClient) -> None:
    # Swagger docs, ReDoc, and OpenAPI schema must be hidden from public discovery
    assert client.get("/docs").status_code == 404
    assert client.get("/redoc").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_api_key_protection_when_configured(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RANDOMOFFICE_API_KEY", "secret-test-key-12345")
    # Without key -> 401 Unauthorized
    resp_unauth = client.post(
        "/api/v1/parse",
        files={"file": ("test.docx", b"PK...", "application/octet-stream")},
    )
    assert resp_unauth.status_code == 401
    assert "invalid or missing API key" in resp_unauth.json()["detail"]

    # With X-API-Key header -> proceeds (rejects with 400 bad format, not 401)
    resp_auth = client.post(
        "/api/v1/parse",
        files={"file": ("test.pdf", b"test", "application/pdf")},
        headers={"X-API-Key": "secret-test-key-12345"},
    )
    assert resp_auth.status_code == 400

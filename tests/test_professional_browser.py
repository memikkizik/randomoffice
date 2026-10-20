"""Desktop-style workflows verified in the shared actual browser workspace."""

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from openpyxl import load_workbook
import pytest

pytest.importorskip('playwright.sync_api')
from test_editor_browser import browser_server, page


@pytest.mark.parametrize('size', [(640, 520), (760, 600), (910, 512), (1366, 768)])
def test_workspace_adapts_to_small_screens(page, size):
    width, height = size
    page.set_viewport_size({'width': width, 'height': height})
    for app in ('word', 'excel', 'powerpoint'):
        page.evaluate('app => activateOfficeApp(app)', app)
        footer = page.locator('.desktop-status-bar').bounding_box()
        assert footer['y'] + footer['height'] <= height + 1
        actions = page.locator('.top-right-cluster').bounding_box()
        assert actions['x'] >= 0 and actions['x'] + actions['width'] <= width + 1
    page.locator('#pill-word').click()
    sheet = page.locator('#word-editor').bounding_box()
    assert sheet['x'] >= 0 and sheet['x'] + sheet['width'] <= width + 1


def test_excel_scroll_and_zoom_work(page):
    page.set_viewport_size({'width': 760, 'height': 600})
    page.evaluate("activateOfficeApp('excel')")
    sizes = page.evaluate("""() => {const pane=document.querySelector('.excel-scroll-area');
        return {client:pane.clientWidth,scroll:pane.scrollWidth};}""")
    assert sizes['client'] <= 760
    assert sizes['scroll'] > sizes['client']
    before = page.locator('#excel-sheet-grid').bounding_box()['width']
    page.evaluate('handleZoomScale(150)')
    assert page.locator('#excel-sheet-grid').bounding_box()['width'] > before * 1.4


def test_word_find_replace_keeps_formatting(page):
    page.evaluate("document.getElementById('word-editor').innerHTML='<p><b>Hello</b> Hello</p>'")
    page.evaluate("openWorkspacePanel('find')")
    page.locator('#workspace-find-input').fill('Hello')
    assert '2 eşleşme' in page.locator('#workspace-find-status').inner_text()
    page.locator('#workspace-replace-input').fill('Merhaba')
    page.evaluate('replaceWorkspaceText(true)')
    assert page.locator('#word-editor b').inner_text() == 'Merhaba'
    assert page.locator('#word-editor').inner_text().count('Merhaba') == 2


def test_excel_cell_styles_survive_switch_and_export(page, tmp_path):
    page.evaluate("createBlankFromHub('excel')")
    page.locator('td[data-row="0"][data-col="0"]').fill('Styled')
    page.evaluate("execOfficeAction('bold');applyColor('foreColor','#107c41')")
    page.evaluate('addNewSheetTab();switchExcelSheet(0)')
    cell = page.locator('td[data-row="0"][data-col="0"]')
    assert cell.evaluate("node=>getComputedStyle(node).fontWeight") == '700'
    with page.expect_download() as download:
        page.evaluate('saveDocument()')
    target = tmp_path / 'styled.xlsx'
    download.value.save_as(target)
    workbook = load_workbook(target)
    assert workbook.worksheets[0]['A1'].font.bold
    assert workbook.worksheets[0]['A1'].font.color.rgb.upper().endswith('107C41')
    workbook.close()


def test_clipboard_grid_preserves_formula_and_leading_zero(page):
    page.evaluate("createBlankFromHub('excel')")
    page.evaluate(r"""handleSheetPaste({clipboardData:{getData:()=> '00123\t2\nItem\t=SUM(B1:B1)'},preventDefault(){}})""")
    assert page.evaluate('excelModel.sheets[0].rows[0][0]') == '00123'
    assert page.evaluate('excelModel.sheets[0].rows[1][1]') == '=SUM(B1:B1)'


def test_page_layout_is_saved_as_docx_section(page, tmp_path):
    page.evaluate("setPageLayout('orientation','landscape');setPageLayout('paper','Letter');setPageLayout('margin','narrow')")
    with page.expect_download() as download:
        page.evaluate('saveDocument()')
    target = tmp_path / 'landscape.docx'
    download.value.save_as(target)
    section = Document(target).sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height
    assert abs(section.left_margin.inches - .5) < .01
    assert 'Letter landscape' in page.locator('#workspace-print-settings').inner_text()


def test_presenter_keeps_media_and_theme(page):
    png='iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aO1sAAAAASUVORK5CYII='
    page.evaluate("""image=>ingestDocumentData({app:'powerpoint',filename:'media.pptx',slides:[
        {title:'Media',paragraphs:['Body'],images:['data:image/png;base64,'+image],tables:[[['A','B']]],theme:'navy'}]})""", png)
    page.evaluate('launchPresenterShow()')
    assert page.locator('#presenter-media img').count() == 1
    assert page.locator('#presenter-media table').inner_text() == 'A\tB'
    assert page.locator('#presenter-modal').get_attribute('data-theme') == 'navy'


def test_desktop_startup_file_is_ingested_once(page):
    calls=page.evaluate("""async()=>{
        let count=0;
        window.pywebview={api:{get_startup_document:async()=>{count++;return {app:'word',filename:'Opened.docx',items:[{type:'paragraph',text:'Opened by Explorer'}]};}}};
        await Promise.all([ingestDesktopStartup(),ingestDesktopStartup()]);return count;
    }""")
    assert calls == 1
    assert page.locator('#word-editor').inner_text().strip() == 'Opened by Explorer'
    assert page.locator('#doc-name-editor').input_value() == 'Opened.docx'


def test_native_drafts_work_when_browser_storage_is_denied(page):
    result=page.evaluate("""async()=>{
        let saved;
        Storage.prototype.setItem=()=>{throw new Error('Access denied');};
        window.pywebview={api:{read_workspace_state:async()=>({state:{app:'word',title:'Draft.docx',wordHtml:'<p>Recovered draft</p>'}}),get_startup_document:async()=>null,
            write_workspace_state:async state=>{saved=state;return {success:true};}}};
        await ingestDesktopStartup();
        document.getElementById('word-editor').innerHTML='<p>Updated draft</p>';persistState();await nativeStateWriteChain;
        return saved;
    }""")
    assert 'Updated draft' in result['wordHtml']
    assert 'Cihazda taslak' in page.locator('#cloud-sync-status').inner_text()


def test_save_and_save_as_use_distinct_native_actions(page):
    result=page.evaluate("""async()=>{
        const calls=[];window.pywebview={api:{
            save_current_file:async()=>{calls.push('current');return {cancelled:true};},
            save_native_file:async()=>{calls.push('as');return {cancelled:true};}}};
        await saveDocument();await saveDocument(true);return calls;
    }""")
    assert result == ['current','as']


def test_removing_prior_slide_keeps_selected_slide(page):
    page.evaluate("""ingestDocumentData({app:'powerpoint',filename:'slides.pptx',slides:[
        {title:'A'},{title:'B'},{title:'C'}]});activeSlideIndex=1;renderPptDeckSidebar();removeSlideFromThumb({stopPropagation(){}},0)""")
    assert page.locator('#ppt-slide-header').inner_text() == 'B'


def test_sorting_moves_styles_with_rows_and_freezes_header(page):
    page.evaluate("""ingestDocumentData({app:'excel',filename:'sort.xlsx',sheets:[{name:'Data',rows:[['Item'],['Zebra'],['Apple']],cell_styles:{'1:0':{bold:true}}}]});sortExcelRows(false);toggleFreezeRow()""")
    assert page.evaluate('excelModel.sheets[0].rows[1][0]') == 'Apple'
    assert page.evaluate("excelModel.sheets[0].cell_styles['2:0'].bold") is True
    assert 'freeze-first-row' in page.locator('#excel-sheet-grid').get_attribute('class')

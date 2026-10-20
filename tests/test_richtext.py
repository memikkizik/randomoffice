import base64
import io
import pytest

from docx import Document
from docx.shared import RGBColor
from openpyxl import load_workbook
from pptx import Presentation

from randomoffice.web_office import export_office_file, open_office_file


PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aO1sAAAAASUVORK5CYII="


def test_word_editor_html_keeps_formatting_lists_and_tables():
    content = export_office_file("word", {"html": '''
        <h1>Heading</h1><p style="text-align:center"><b>Bold</b> <i>Italic</i>
        <span style="color:#107c41;font-size:18pt">Green</span></p>
        <ol><li>First</li><li>Second</li></ol>
        <table><tr><td><b>Cell</b></td><td>Other</td></tr></table>
    '''})
    document = Document(io.BytesIO(content))
    runs = [run for p in document.paragraphs for run in p.runs]
    assert any(r.bold and r.text == "Bold" for r in runs)
    assert any(r.italic and r.text == "Italic" for r in runs)
    assert any(r.font.color.rgb == RGBColor.from_string("107C41") and r.font.size.pt == 18 for r in runs if r.text == "Green")
    assert sum(p.style.name == "List Number" for p in document.paragraphs) == 2
    assert document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold


def test_word_open_exposes_safe_formatting(tmp_path):
    source = tmp_path / "styled.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("<unsafe>").bold = True
    document.save(source)
    data = open_office_file(source)
    assert "font-weight:bold" in data["items"][0]["html"]
    assert "&lt;unsafe&gt;" in data["items"][0]["html"]
    reopened = Document(io.BytesIO(export_office_file("word", data)))
    assert reopened.paragraphs[0].runs[0].bold


def test_word_embedded_image_survives(tmp_path):
    content = export_office_file("word", {"html": f'<p>Image<img src="data:image/png;base64,{PNG}"></p>'})
    document = Document(io.BytesIO(content))
    assert len(document.inline_shapes) == 1
    source = tmp_path / "image.docx"
    source.write_bytes(content)
    assert "data:image/png;base64," in open_office_file(source)["items"][0]["html"]


def test_word_links_survive(tmp_path):
    content = export_office_file("word", {"html": '<p><a href="https://example.com">Example</a></p>'})
    source = tmp_path / "links.docx"
    source.write_bytes(content)
    assert 'href="https://example.com"' in open_office_file(source)["items"][0]["html"]


def test_export_does_not_fetch_external_images():
    content = export_office_file("word", {"html": '<p>Safe<img src="http://127.0.0.1:1/secret"></p><script>bad()</script>'})
    document = Document(io.BytesIO(content))
    assert len(document.inline_shapes) == 0
    assert document.paragraphs[0].text == "Safe"


def test_excel_number_format_keeps_numeric_type():
    content = export_office_file("excel", {"sheets": [{"name": "Money", "rows": [[1234.5]], "number_formats": {"0:0": '"₺" #,##0.00'}}]})
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.active["A1"].value == 1234.5
    assert "₺" in workbook.active["A1"].number_format
    workbook.close()


def test_ppt_embedded_image_and_theme_survive(tmp_path):
    content = export_office_file("slides", {"slides": [{"title": "Image", "images": [f"data:image/png;base64,{PNG}"], "theme": "navy"}]})
    presentation = Presentation(io.BytesIO(content))
    assert any(shape.shape_type == 13 for shape in presentation.slides[0].shapes)
    assert str(presentation.slides[0].background.fill.fore_color.rgb) == "0F2B5C"
    source = tmp_path / "picture.pptx"
    source.write_bytes(content)
    assert open_office_file(source)["slides"][0]["images"]


@pytest.mark.parametrize(("kind", "data"), [
    ("word", {"items": [None]}),
    ("excel", {"sheets": [{"name": "Invalid/Name", "rows": []}]}),
    ("excel", {"sheets": [{"name": "Data"}, {"name": "data"}]}),
    ("excel", {"sheets": [{"rows": [None]}]}),
    ("powerpoint", {"slides": [{"paragraphs": [None]}]}),
    ("powerpoint", {"slides": [{"tables": None}]}),
])
def test_invalid_editor_payload_is_rejected(kind, data):
    with pytest.raises(ValueError):
        export_office_file(kind, data)

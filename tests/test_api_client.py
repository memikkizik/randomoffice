"""Tests for examples/api_client.py covering health, parse, convert, and roundtrip endpoints."""

from pathlib import Path
import socket
import sys
import threading
import time

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pytest
import uvicorn

# Ensure project root in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from examples.api_client import request_json, roundtrip_document, upload
from scripts.generate_samples import generate_docx, generate_xlsx, generate_pptx


@pytest.fixture(scope="module")
def api_server_url():
    from randomoffice.api.app import create_app

    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        listener.listen(128)
        port = listener.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(create_app(), log_level="error"))
        thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
        thread.start()
        try:
            deadline = time.monotonic() + 10
            while not server.started:
                if not thread.is_alive() or time.monotonic() > deadline:
                    raise RuntimeError("API test server failed to start")
                time.sleep(0.02)
            yield f"http://127.0.0.1:{port}"
        finally:
            server.should_exit = True
            thread.join(timeout=10)


def test_api_client_health(api_server_url: str):
    data = request_json(f"{api_server_url}/health")
    assert data["status"] == "healthy"
    assert "uptime_seconds" in data


def test_api_client_roundtrip_docx(api_server_url: str, tmp_path: Path):
    source_dir = tmp_path / "src"
    source_dir.mkdir()
    source = generate_docx(source_dir)

    target = tmp_path / "out" / "roundtrip.docx"
    result = roundtrip_document(api_server_url, source, target)
    assert result.exists()
    assert result.stat().st_size > 0

    doc = Document(result)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "RandomOffice Örnek Türkçe Belge" in text or "Kalın metin" in text


def test_api_client_roundtrip_xlsx(api_server_url: str, tmp_path: Path):
    source_dir = tmp_path / "src"
    source_dir.mkdir(exist_ok=True)
    source = generate_xlsx(source_dir)

    target = tmp_path / "out" / "roundtrip.xlsx"
    result = roundtrip_document(api_server_url, source, target)
    assert result.exists()

    wb = load_workbook(result)
    assert len(wb.sheetnames) >= 1
    sheet = wb.active
    # Verify values and formulas preserved
    assert sheet["A3"].value == "00123"
    assert sheet["C6"].value == "=SUM(C3:C5)"
    wb.close()


def test_api_client_roundtrip_pptx(api_server_url: str, tmp_path: Path):
    source_dir = tmp_path / "src"
    source_dir.mkdir(exist_ok=True)
    source = generate_pptx(source_dir)

    target = tmp_path / "out" / "roundtrip.pptx"
    result = roundtrip_document(api_server_url, source, target)
    assert result.exists()

    prs = Presentation(result)
    assert len(prs.slides) >= 1
    slide = prs.slides[0]
    assert slide.has_notes_slide
    assert "hatırlatma" in slide.notes_slide.notes_text_frame.text


def test_api_client_same_source_target_rejected(api_server_url: str, tmp_path: Path):
    dummy = tmp_path / "test.docx"
    dummy.write_text("dummy")
    with pytest.raises(ValueError, match="aynı olamaz"):
        roundtrip_document(api_server_url, dummy, dummy)


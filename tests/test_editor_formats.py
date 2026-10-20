"""Round-trip checks for editor styles in real Office archives."""

import base64
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Mm
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from randomoffice.web_office import export_office_file, open_office_file


PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aO1sAAAAASUVORK5CYII="


def test_excel_styles_formulas_and_freeze_survive_two_exports(tmp_path):
    style = {"bold": True, "italic": True, "underline": True, "strike": True,
             "color": "#185ABD", "background": "#FFF100", "font_name": "Arial",
             "font_size": 16, "alignment": "center"}
    payload = {"sheets": [{"name": "Bütçe", "rows": [[1234.5, "=A1*2"]],
                            "cell_styles": {"0:0": style}, "freeze_first_row": True,
                            "number_formats": {"0:0": '"₺" #,##0.00'}}]}
    source = tmp_path / "styled.xlsx"
    source.write_bytes(export_office_file("excel", payload))
    parsed = open_office_file(source)
    sheet = parsed["sheets"][0]
    assert sheet["cell_styles"]["0:0"] == style
    assert sheet["freeze_first_row"] is True
    assert sheet["number_formats"]["0:0"] == '"₺" #,##0.00'
    workbook = load_workbook(io.BytesIO(export_office_file("excel", parsed)), data_only=False)
    try:
        cell = workbook.active["A1"]
        assert cell.value == 1234.5
        assert workbook.active["B1"].value == "=A1*2"
        assert cell.font.bold and cell.font.italic and cell.font.strike
        assert cell.font.underline == "single"
        assert cell.font.color.rgb[-6:] == "185ABD"
        assert cell.fill.fgColor.rgb[-6:] == "FFF100"
        assert cell.alignment.horizontal == "center"
        assert cell.font.name == "Arial" and cell.font.sz == 16
        assert workbook.active.freeze_panes == "A2"
    finally:
        workbook.close()


def test_excel_empty_styled_cell_and_explicit_false_styles_survive(tmp_path):
    source = tmp_path / "blank-styles.xlsx"
    source.write_bytes(export_office_file("excel", {"sheets": [{"name": "Data", "rows": [["Value"]],
        "cell_styles": {"2:2": {"background": "#12AB34", "bold": False, "underline": False}},
        "freeze_first_row": False}]}))
    data = open_office_file(source)["sheets"][0]
    assert data["cell_styles"]["2:2"]["background"] == "#12AB34"
    assert data["cell_styles"]["2:2"]["bold"] is False
    assert data["rows"][2][2] is None
    assert data["freeze_first_row"] is False


def test_excel_import_resolves_indexed_and_theme_colors(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Color, Font, PatternFill

    source = tmp_path / "theme.xlsx"
    workbook = Workbook()
    cell = workbook.active["A1"]
    cell.value = "Theme"
    cell.font = Font(color=Color(indexed=10))
    cell.fill = PatternFill(fill_type="solid", fgColor=Color(theme=4))
    workbook.save(source)
    style = open_office_file(source)["sheets"][0]["cell_styles"]["0:0"]
    assert style["color"] == "#FF0000"
    assert style["background"] == "#4F81BD"


@pytest.mark.parametrize(("layout", "width", "height", "margins"), [
    ({"orientation": "landscape", "paper": "A4", "margin": "narrow"}, Mm(297), Mm(210), (.5, .5, .5, .5)),
    ({"orientation": "portrait", "paper": "Letter", "margin": "wide"}, Inches(8.5), Inches(11), (1, 1, 2, 2)),
    ({"orientation": "portrait", "paper": "A4", "margin": "normal"}, Mm(210), Mm(297), (1, 1, 1, 1)),
])
def test_word_page_layout_is_real_section_metadata_and_roundtrips(tmp_path, layout, width, height, margins):
    source = tmp_path / "layout.docx"
    source.write_bytes(export_office_file("word", {"title": "Layout", "html": "<p><b>Styled content</b></p>", "page_layout": layout}))
    document = Document(source)
    section = document.sections[0]
    # OOXML stores dimensions in twips; conversion can round by one twip.
    assert abs(section.page_width-width) <= 635
    assert abs(section.page_height-height) <= 635
    assert section.orientation == (WD_ORIENT.LANDSCAPE if layout["orientation"] == "landscape" else WD_ORIENT.PORTRAIT)
    assert [section.top_margin, section.bottom_margin, section.left_margin, section.right_margin] == [Inches(value) for value in margins]
    parsed = open_office_file(source)
    assert parsed["page_layout"] == layout
    reopened = Document(io.BytesIO(export_office_file("word", parsed)))
    assert abs(reopened.sections[0].page_width-width) <= 635
    assert reopened.paragraphs[0].runs[0].bold


def test_word_item_export_also_applies_layout():
    document = Document(io.BytesIO(export_office_file("word", {
        "page_layout": {"orientation": "landscape", "paper": "Letter", "margin": "wide"},
        "items": [{"type": "paragraph", "text": "Plain"}]})))
    assert document.sections[0].page_width == Inches(11)
    assert document.paragraphs[0].text == "Plain"


def test_ppt_rich_runs_theme_images_and_tables_survive_roundtrip(tmp_path):
    payload = {"slides": [{"title": "Fallback", "paragraphs": ["Fallback body"],
        "title_html": '<strong>Title &lt;safe&gt;</strong>',
        "body_html": '<div>• <b>Bold</b> and <i>italic</i></div><div>• <span style="color:#E81123;font-size:24pt;font-family:Arial"><u>Red</u></span><br>Next line</div>',
        "theme": "navy", "notes": "Private notes", "images": [f"data:image/png;base64,{PNG}"],
        "tables": [[["Field", "Value"], ["Status", "Ready"]]]}]}
    source = tmp_path / "rich.pptx"
    source.write_bytes(export_office_file("powerpoint", payload))
    presentation = Presentation(source)
    slide = presentation.slides[0]
    assert slide.shapes.title.text == "Title <safe>"
    assert slide.shapes.title.text_frame.paragraphs[0].runs[0].font.bold
    body = slide.placeholders[1].text_frame
    assert [p.text for p in body.paragraphs] == ["Bold and italic", "Red\vNext line"]
    assert body.paragraphs[0].runs[0].font.bold
    red = body.paragraphs[1].runs[0]
    assert str(red.font.color.rgb) == "E81123"  # Explicit color survives dark theme.
    assert red.font.size.pt == 24 and red.font.underline
    assert red.font.name == "Arial"
    parsed = open_office_file(source)
    record = parsed["slides"][0]
    assert record["theme"] == "navy"
    assert record["notes"] == "Private notes"
    assert "font-weight:bold" in record["title_html"]
    assert "color:#E81123" in record["body_html"]
    assert "<br>" in record["body_html"]
    assert "&lt;safe&gt;" in record["title_html"]
    assert base64.b64decode(record["images"][0].split(",", 1)[1]) == base64.b64decode(PNG)
    reopened = Presentation(io.BytesIO(export_office_file("slides", parsed)))
    assert reopened.slides[0].placeholders[1].text_frame.paragraphs[1].runs[0].font.underline
    assert any(shape.has_table for shape in reopened.slides[0].shapes)
    assert next(shape.image.blob for shape in reopened.slides[0].shapes if shape.shape_type == 13) == base64.b64decode(PNG)


def test_ppt_html_rejects_active_content_and_never_fetches_external_resources(tmp_path):
    source = tmp_path / "safe.pptx"
    source.write_bytes(export_office_file("slides", {"slides": [{
        "title_html": '<b onclick="bad()">Safe</b><script>bad()</script>',
        "body_html": '<p>Body<img src="http://127.0.0.1:1/secret"><iframe src="http://127.0.0.1:1/other">bad</iframe><span style="background:url(http://127.0.0.1:1)"> text</span></p>'}]}))
    parsed = open_office_file(source)["slides"][0]
    assert parsed["title"] == "Safe"
    assert parsed["paragraphs"] == ["Body text"]
    assert parsed["images"] == []
    assert all(value not in parsed["title_html"]+parsed["body_html"] for value in ("onclick", "script", "iframe", "url(", "http://"))


def test_ppt_body_matching_title_is_not_discarded(tmp_path):
    source = tmp_path / "repeated.pptx"
    source.write_bytes(export_office_file("slides", {"slides": [{"title": "Repeated", "paragraphs": ["Repeated"]}]}))
    record = open_office_file(source)["slides"][0]
    assert record["paragraphs"] == ["Repeated"]
    assert "Repeated" in record["body_html"]


@pytest.mark.parametrize(("kind", "data", "message"), [
    ("word", {"page_layout": {"paper": "Huge"}}, "page_layout paper"),
    ("word", {"page_layout": []}, "page_layout"),
    ("excel", {"sheets": [{"cell_styles": {"-1:0": {"bold": True}}}]}, "addresses"),
    ("excel", {"sheets": [{"cell_styles": {"0:0": {"color": "url(http://evil)"}}}]}, "#RRGGBB"),
    ("excel", {"sheets": [{"cell_styles": {"0:0": {"font_size": float("nan")}}}]}, "font_size"),
    ("excel", {"sheets": [{"freeze_first_row": "yes"}]}, "freeze_first_row"),
    ("slides", {"slides": [{"body_html": []}]}, "body_html"),
])
def test_malformed_formatting_reports_clear_error(kind, data, message):
    with pytest.raises(ValueError, match=message):
        export_office_file(kind, data)

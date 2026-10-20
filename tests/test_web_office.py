import io
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
import openpyxl
from openpyxl import Workbook
from pptx import Presentation
import pytest

from randomoffice.api.app import app
from randomoffice.web_office import export_office_file, open_office_file


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_web_office_open_and_export_docx(tmp_path: Path) -> None:
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Web Office Title", level=1)
    doc.add_paragraph("Paragraph body")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Col1"
    tbl.rows[0].cells[1].text = "Col2"
    tbl.rows[1].cells[0].text = "Val1"
    tbl.rows[1].cells[1].text = "Val2"
    doc.save(str(doc_path))

    data = open_office_file(doc_path)
    assert data["app"] == "word"
    assert data["format"] == ".docx"
    assert len(data["items"]) >= 2

    # Test export
    exported_bytes = export_office_file("word", data)
    assert len(exported_bytes) > 1000

    # Verify exported docx can be opened again
    reopened = Document(io.BytesIO(exported_bytes))
    assert len(reopened.paragraphs) >= 1


def test_web_office_open_and_export_xlsx(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty"])
    ws.append(["Pen", 10])
    wb.save(str(xlsx_path))

    data = open_office_file(xlsx_path)
    assert data["app"] == "excel"
    assert data["sheets"][0]["name"] == "Sales"
    assert data["sheets"][0]["rows"][0][0] == "Item"

    exported_bytes = export_office_file("excel", data)
    assert len(exported_bytes) > 1000

    reopened = openpyxl.load_workbook(io.BytesIO(exported_bytes), data_only=False)
    assert reopened["Sales"]["B2"].value == 10
    reopened.close()


def test_xlsx_formulas_and_leading_zero_values_survive_round_trip(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "00123"
    worksheet["B1"] = "=SUM(1,2)"
    workbook.save(source)

    data = open_office_file(source)
    exported = export_office_file("excel", data)
    reopened = openpyxl.load_workbook(io.BytesIO(exported), data_only=False)

    assert reopened.active["A1"].value == "00123"
    assert reopened.active["B1"].value == "=SUM(1,2)"
    reopened.close()


def test_web_office_open_and_export_pptx(tmp_path: Path) -> None:
    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Presentation Slide"
    prs.save(str(pptx_path))

    data = open_office_file(pptx_path)
    assert data["app"] == "powerpoint"
    assert data["slides"][0]["title"] == "Presentation Slide"

    exported_bytes = export_office_file("powerpoint", data)
    assert len(exported_bytes) > 1000


def test_pptx_tables_survive_round_trip(tmp_path: Path) -> None:
    source = tmp_path / "table-deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, 0, 0, 100, 100).table
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Ready"
    presentation.save(source)

    data = open_office_file(source)
    exported = export_office_file("powerpoint", data)
    reopened = Presentation(io.BytesIO(exported))

    assert any(shape.has_table for shape in reopened.slides[0].shapes)


def test_api_serve_html_index(client: TestClient) -> None:
    response = client.get("/")
    assert response.status_code == 200
    assert "RandomOffice" in response.text
    assert "Word" in response.text
    assert "Excel" in response.text


def test_api_office_open_and_export(client: TestClient) -> None:
    doc = Document()
    doc.add_paragraph("API Open Test")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # Open via API
    res_open = client.post(
        "/api/v1/office/open",
        files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res_open.status_code == 200
    data = res_open.json()
    assert data["app"] == "word"

    # Export via API
    res_export = client.post(
        "/api/v1/office/export?doc_type=word&filename=test.docx",
        json=data,
    )
    assert res_export.status_code == 200
    assert len(res_export.content) > 1000


def test_xlsx_cells_outside_original_viewport_survive(tmp_path: Path) -> None:
    source = tmp_path / "large.xlsx"
    workbook = Workbook()
    workbook.active["AA101"] = "Keep this value"
    workbook.save(source)
    data = open_office_file(source)
    reopened = openpyxl.load_workbook(io.BytesIO(export_office_file("excel", data)))
    assert reopened.active["AA101"].value == "Keep this value"
    reopened.close()


def test_pptx_speaker_notes_survive(tmp_path: Path) -> None:
    exported = export_office_file("slides", {
        "slides": [{"title": "Notes", "notes": "Speaker instructions", "paragraphs": []}]
    })
    source = tmp_path / "notes.pptx"
    source.write_bytes(exported)
    assert open_office_file(source)["slides"][0]["notes"] == "Speaker instructions"


def test_api_export_supports_turkish_download_names(client: TestClient) -> None:
    response = client.post(
        "/api/v1/office/export",
        params={"doc_type": "word", "filename": "?al??ma raporu.docx"},
        json={"items": []},
    )
    assert response.status_code == 200
    from urllib.parse import unquote
    disposition = response.headers["content-disposition"]
    assert unquote(disposition.split("UTF-8''", 1)[1]) == "?al??ma raporu.docx"


def test_api_rejects_invalid_xlsx_archive(client: TestClient) -> None:
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("unrelated.txt", "Not a workbook")
    response = client.post(
        "/api/v1/office/open",
        files={"file": ("broken.xlsx", buffer.getvalue())},
    )
    assert response.status_code == 422


def test_xlsx_rejects_oversized_grid_without_truncation(tmp_path: Path) -> None:
    from randomoffice.errors import DocumentParseError
    source = tmp_path / "oversized.xlsx"
    workbook = Workbook()
    workbook.active["Z10000"] = "Far away"
    workbook.save(source)
    with pytest.raises(DocumentParseError, match="200,000 cells"):
        open_office_file(source)

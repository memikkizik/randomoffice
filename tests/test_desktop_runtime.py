"""Native file identity, safe writes, and startup contracts without a visible GUI."""

import base64
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import os
import runpy
import sys
from types import SimpleNamespace

from docx import Document
import pytest

import randomoffice.desktop as desktop
from randomoffice.__main__ import main


class EventHandlers:
    def __init__(self):
        self.handlers = []

    def __iadd__(self, handler):
        self.handlers.append(handler)
        return self


class FakeWindow:
    def __init__(self, selections=(), confirmations=()):
        self.selections = iter(selections)
        self.confirmations = iter(confirmations)
        self.close_dialogs = []
        self.dialogs = []
        self.scripts = []
        self.events = SimpleNamespace(loaded=EventHandlers(), closing=EventHandlers())

    def create_file_dialog(self, kind, **options):
        self.dialogs.append((kind, options))
        return next(self.selections, None)

    def evaluate_js(self, script):
        self.scripts.append(script)

    def create_confirmation_dialog(self, title, message):
        self.close_dialogs.append((title, message))
        return next(self.confirmations, False)


def make_document(path: Path, text="Original content") -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


@pytest.mark.parametrize("area", [
    desktop.DesktopWorkArea(800, 600),
    desktop.DesktopWorkArea(1024, 640),
    desktop.DesktopWorkArea(640, 400),
    desktop.DesktopWorkArea(2560, 1400, -2560, -1400),
])
def test_restored_window_fits_monitor_work_area(area):
    options = desktop.adaptive_window_options(area)
    assert options["maximized"] is True
    assert area.x <= options["x"] <= area.x + area.width - options["width"]
    assert area.y <= options["y"] <= area.y + area.height - options["height"]
    assert 0 < options["min_size"][0] <= options["width"] <= area.width
    assert 0 < options["min_size"][1] <= options["height"] <= area.height


def test_user_storage_is_stable_outside_frozen_bundle(tmp_path, monkeypatch):
    monkeypatch.setattr(desktop.sys, "platform", "win32")
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "local"))
    monkeypatch.delenv("RANDOMOFFICE_DATA_DIR", raising=False)
    expected = tmp_path / "local" / "RandomOffice"
    monkeypatch.setattr(desktop.sys, "_MEIPASS", str(tmp_path / "extracted-a"), raising=False)
    assert desktop.get_user_data_directory() == expected
    monkeypatch.setattr(desktop.sys, "_MEIPASS", str(tmp_path / "extracted-b"), raising=False)
    assert desktop.get_user_data_directory() == expected
    assert not expected.exists()  # Inspection must not create a user profile.


def test_user_storage_can_be_isolated_without_touching_real_profile(tmp_path, monkeypatch):
    monkeypatch.setenv("RANDOMOFFICE_DATA_DIR", str(tmp_path / "environment"))
    assert desktop.get_user_data_directory() == tmp_path / "environment"
    assert desktop.get_user_data_directory(tmp_path / "explicit") == tmp_path / "explicit"


def test_startup_file_is_consumed_once_and_path_never_embedded_in_javascript(tmp_path):
    source = make_document(tmp_path / "rapor '); alert(1); (' İstanbul.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path / "profile")
    window = FakeWindow()
    bridge.set_window(window)
    assert bridge.get_desktop_state()["startup_pending"] is True
    bridge.notify_editor_ready()
    assert "ingestDesktopStartup" in window.scripts[0]
    assert source.name not in window.scripts[0]
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: bridge.get_startup_document(), range(2)))
    received = [result for result in results if result is not None]
    assert len(received) == 1
    assert received[0]["filename"] == source.name
    assert received[0]["items"][0]["text"] == "Original content"
    state = bridge.get_desktop_state()
    assert state["startup_pending"] is False
    assert state["current_files"]["word"]["saved_path"] == str(source)


@pytest.mark.parametrize("filename", ["missing.docx", "untrusted.exe"])
def test_bad_startup_path_returns_an_error_without_crashing_or_retrying(tmp_path, filename):
    bridge = desktop.NativeDesktopBridge(tmp_path / filename, data_directory=tmp_path)
    assert "error" in bridge.get_startup_document()
    assert bridge.get_startup_document() is None
    assert bridge.get_desktop_state()["current_files"] == {}


def test_save_overwrites_only_current_native_source_and_keeps_filename_server_side(tmp_path):
    source = make_document(tmp_path / "rapor.docx")
    other = make_document(tmp_path / "protected.docx", "Protected")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    window = FakeWindow()
    bridge.set_window(window)
    saved = bridge.save_current_file("word", str(other), {"html": "<p>Changed</p>", "saved_path": str(other)})
    assert saved["success"] is True
    assert saved["filename"] == source.name
    assert saved["saved_path"] == str(source)
    assert saved["desktop_file"]["display_name"] == source.name
    assert Document(source).paragraphs[0].text == "Changed"
    assert Document(other).paragraphs[0].text == "Protected"
    assert window.dialogs == []


def test_save_as_then_ctrl_s_reuses_destination_and_does_not_reopen_dialog(tmp_path):
    destination = tmp_path / "İş planı.docx"
    window = FakeWindow([str(destination)])
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    bridge.set_window(window)
    first = bridge.save_current_file("word", "İş planı", {"html": "<p>First</p>"})
    assert first["success"] is True
    assert window.dialogs[0][1]["save_filename"] == "İş planı.docx"
    second = bridge.save_current_file("docx", "Changed suggestion.docx", {"html": "<p>Second</p>"})
    assert second["success"] is True
    assert Document(destination).paragraphs[0].text == "Second"
    assert len(window.dialogs) == 1


def test_new_document_clears_identity_and_cancelled_save_as_preserves_old_document(tmp_path):
    source = make_document(tmp_path / "original.docx")
    before = source.read_bytes()
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    window = FakeWindow([None, None])
    bridge.set_window(window)
    assert bridge.save_native_file("word", "copy.docx", {})["cancelled"] is True
    assert bridge.get_desktop_state()["current_files"]["word"]["saved_path"] == str(source)
    assert bridge.reset_current_file("word")["success"] is True
    assert bridge.save_current_file("word", "new.docx", {"html": "<p>New</p>"})["cancelled"] is True
    assert source.read_bytes() == before
    assert bridge.get_desktop_state()["current_files"] == {}


def test_export_failure_leaves_original_document_intact(tmp_path):
    source = make_document(tmp_path / "original.docx")
    before = source.read_bytes()
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    result = bridge.save_current_file("word", source.name, {"html": 123})
    assert result["success"] is False
    assert "error" in result
    assert source.read_bytes() == before
    assert list(tmp_path.glob("*.tmp")) == []


def test_atomic_replacement_failure_preserves_original_and_cleans_temporary_file(tmp_path, monkeypatch):
    doc_dir = tmp_path / "docs"
    doc_dir.mkdir()
    source = make_document(doc_dir / "original.docx")
    before = source.read_bytes()
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path / "data")
    bridge.get_startup_document()

    def locked_file(*_):
        raise PermissionError("The document is open in another application")

    monkeypatch.setattr(desktop.os, "replace", locked_file)
    result = bridge.save_current_file("word", source.name, {"html": "<p>Changed</p>"})
    assert result["success"] is False
    assert "another application" in result["error"]
    assert source.read_bytes() == before
    assert list(doc_dir.iterdir()) == [source]


def test_dropped_document_cannot_overwrite_previous_native_file(tmp_path):
    source = make_document(tmp_path / "original.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    result = bridge.open_file_data("dropped.docx", base64.b64encode(source.read_bytes()).decode("ascii"))
    assert result["filename"] == "dropped.docx"
    assert bridge.get_desktop_state()["current_files"] == {}


def test_open_native_tracks_successful_selection_and_cancel_preserves_it(tmp_path):
    source = make_document(tmp_path / "selected.docx")
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    window = FakeWindow([(str(source),), None])
    bridge.set_window(window)
    opened = bridge.open_native_file()
    assert opened["saved_path"] == str(source)
    assert bridge.open_native_file() is None
    assert bridge.get_desktop_state()["current_files"]["word"]["filename"] == source.name


def test_save_as_correct_extension_added_and_wrong_extension_rejected(tmp_path):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    window = FakeWindow([str(tmp_path / "book"), str(tmp_path / "mismatch.docx")])
    bridge.set_window(window)
    result = bridge.save_native_file("excel", "book.xlsx", {})
    assert result["success"] is True
    assert (tmp_path / "book.xlsx").is_file()
    rejected = bridge.save_native_file("excel", "book.xlsx", {})
    assert rejected["success"] is False
    assert not (tmp_path / "mismatch.docx").exists()


def test_per_application_current_files_remain_independent(tmp_path):
    source = make_document(tmp_path / "word.docx")
    workbook = tmp_path / "excel.xlsx"
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    bridge.set_window(FakeWindow([str(workbook)]))
    assert bridge.save_current_file("excel", "excel.xlsx", {})["success"] is True
    assert set(bridge.get_desktop_state()["current_files"]) == {"word", "excel"}
    bridge.reset_current_file("excel")
    assert set(bridge.get_desktop_state()["current_files"]) == {"word"}


def test_runtime_starts_persistent_maximized_window_and_hooks_ready_event(tmp_path, monkeypatch):
    calls = {}
    window = FakeWindow()

    def create_window(**options):
        calls["window"] = options
        return window

    def start(**options):
        calls["start"] = options
        for handler in window.events.loaded.handlers:
            handler()

    monkeypatch.setattr(desktop, "enable_dpi_awareness", lambda: None)
    monkeypatch.setattr(desktop, "detect_work_area", lambda: desktop.DesktopWorkArea(960, 600))
    monkeypatch.setattr(desktop, "render_office_ui", lambda *_: "<html>Ready</html>")
    monkeypatch.setattr(desktop.webview, "create_window", create_window)
    monkeypatch.setattr(desktop.webview, "start", start)
    desktop.run_desktop_app(tmp_path / "not-consumed-until-ready.docx", data_directory=tmp_path / "data")
    options = calls["window"]
    assert options["maximized"] is True
    assert options["confirm_close"] is False
    assert len(window.events.closing.handlers) == 1
    assert options["width"] <= 960
    assert options["height"] <= 600
    assert options["min_size"] == (640, 420)
    assert options["js_api"].get_desktop_state()["startup_pending"] is True
    assert window.scripts and "ingestDesktopStartup" in window.scripts[0]
    assert calls["start"]["private_mode"] is False
    assert calls["start"]["storage_path"] == str(tmp_path / "data" / "WebView2")
    assert (tmp_path / "data" / "WebView2").is_dir()


def test_crash_log_appends_to_user_data_instead_of_executable_directory(tmp_path):
    first = desktop.write_crash_log("First startup failure", tmp_path / "data")
    second = desktop.write_crash_log("Second startup failure", tmp_path / "data")
    assert first == second == tmp_path / "data" / "logs" / "randomoffice-crash.log"
    assert "First startup failure" in first.read_text(encoding="utf-8")
    assert "Second startup failure" in first.read_text(encoding="utf-8")


def test_native_workspace_state_survives_restart_without_browser_local_storage(tmp_path):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path / "data")
    assert bridge.read_workspace_state() == {"success": True, "state": {}}
    assert not (tmp_path / "data").exists()
    state = {"theme": "dark", "word-draft": {"html": "<p>İş planı</p>"}, "recent": ["Çalışma.xlsx"]}
    assert bridge.write_workspace_state(state) == {"success": True}
    restarted = desktop.NativeDesktopBridge(data_directory=tmp_path / "data")
    assert restarted.read_workspace_state() == {"success": True, "state": state}


def test_invalid_native_workspace_state_preserves_previous_saved_state(tmp_path, monkeypatch):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    state = {"word-draft": "Original"}
    assert bridge.write_workspace_state(state)["success"] is True
    assert bridge.write_workspace_state(["invalid"])["success"] is False
    assert bridge.write_workspace_state({"not-json": float("nan")})["success"] is False
    monkeypatch.setattr(desktop, "_MAX_WORKSPACE_STATE_BYTES", 30)
    assert bridge.write_workspace_state({"draft": "x" * 100})["success"] is False
    assert bridge.read_workspace_state() == {"success": True, "state": state}


def test_corrupt_native_workspace_is_reported_without_deleting_recovery_file(tmp_path):
    path = tmp_path / "workspace-state.json"
    path.write_text("{broken", encoding="utf-8")
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    assert bridge.read_workspace_state()["success"] is False
    assert path.read_text(encoding="utf-8") == "{broken"


def test_native_workspace_failed_disk_write_preserves_previous_draft(tmp_path, monkeypatch):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    prior = {"wordHtml": "<p>Saved draft</p>"}
    assert bridge.write_workspace_state(prior)["success"] is True

    def locked_file(*_):
        raise PermissionError("Workspace locked")

    monkeypatch.setattr(desktop.os, "replace", locked_file)
    result = bridge.write_workspace_state({"wordHtml": "<p>New draft</p>"})
    assert result["success"] is False
    assert "Workspace locked" in result["error"]
    assert bridge.read_workspace_state()["state"] == prior
    assert [path.name for path in tmp_path.iterdir()] == ["workspace-state.json"]


def test_desktop_cli_passes_optional_document_path_without_internal_port(tmp_path, monkeypatch):
    calls = []
    path = tmp_path / "Çalışma planı.docx"
    monkeypatch.setattr(desktop, "run_desktop_app", lambda **options: calls.append(options))
    assert main(["desktop", str(path)]) == 0
    assert calls == [{"startup_document": path}]


def test_frozen_entrypoint_forwards_double_click_document_argument(tmp_path, monkeypatch):
    calls = []
    path = tmp_path / "Çalışma planı.docx"
    monkeypatch.setattr(desktop, "run_desktop_app", lambda **options: calls.append(options))
    launcher = Path(__file__).resolve().parents[1] / "run_desktop.py"
    monkeypatch.setattr(sys, "argv", [str(launcher), str(path)])
    runpy.run_path(str(launcher), run_name="__main__")
    assert calls == [{"startup_document": path}]


def test_native_save_rejects_external_edits_without_overwriting_them(tmp_path):
    source = make_document(tmp_path / "shared.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    make_document(source, "Changed in Microsoft Word")
    external_bytes = source.read_bytes()
    result = bridge.save_current_file("word", source.name, {"html": "<p>My local edits</p>"})
    assert result["success"] is False
    assert result["conflict"] is True
    assert "başka bir uygulamada" in result["error"]
    assert source.read_bytes() == external_bytes


def test_native_save_rejects_same_size_and_timestamp_external_edits(tmp_path):
    source = make_document(tmp_path / "same-size.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    previous_stat = source.stat()
    changed = bytearray(source.read_bytes())
    changed[-1] ^= 1
    source.write_bytes(changed)
    os.utime(source, ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns))
    result = bridge.save_current_file("word", source.name, {"html": "<p>Local edits</p>"})
    assert result["conflict"] is True
    assert source.read_bytes() == changed


def test_native_save_rejects_deleted_source_without_recreating_it(tmp_path):
    source = make_document(tmp_path / "deleted.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    source.unlink()
    result = bridge.save_current_file("word", source.name, {"html": "<p>Local edits</p>"})
    assert result["conflict"] is True
    assert not source.exists()


def test_save_as_cannot_bypass_conflict_by_selecting_the_same_source(tmp_path):
    source = make_document(tmp_path / "shared.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()
    bridge.set_window(FakeWindow([str(source)]))
    make_document(source, "External version")
    result = bridge.save_native_file("word", source.name, {"html": "<p>My edits</p>"})
    assert result["conflict"] is True
    assert Document(source).paragraphs[0].text == "External version"


def test_second_conflict_check_catches_a_change_during_export(tmp_path, monkeypatch):
    source = make_document(tmp_path / "shared.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    bridge.get_startup_document()

    def concurrent_editor(*_):
        make_document(source, "Changed while exporting")
        return b"replacement must not reach disk"

    monkeypatch.setattr(bridge, "_export_preserving_source", concurrent_editor)
    result = bridge.save_current_file("word", source.name, {"html": "<p>Local edits</p>"})
    assert result["conflict"] is True
    assert Document(source).paragraphs[0].text == "Changed while exporting"
    assert list(tmp_path.glob("*.tmp")) == []


def test_current_file_metadata_reports_read_only_and_native_fingerprint(tmp_path, monkeypatch):
    source = make_document(tmp_path / "read-only.docx")
    bridge = desktop.NativeDesktopBridge(source, data_directory=tmp_path)
    opened = bridge.get_startup_document()
    fingerprint = opened["desktop_file"]["fingerprint"]
    assert len(fingerprint["sha256"]) == 64
    assert fingerprint["size"] == source.stat().st_size
    monkeypatch.setattr(desktop.os, "access", lambda *_: False)
    metadata = bridge.get_desktop_state()["current_files"]["word"]
    assert metadata["read_only"] is True
    assert metadata["fingerprint"] == fingerprint
    metadata["fingerprint"]["sha256"] = "UI cannot change native identity"
    assert bridge.get_desktop_state()["current_files"]["word"]["fingerprint"] == fingerprint


def test_clean_native_close_never_shows_confirmation(tmp_path):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    window = FakeWindow()
    bridge.set_window(window)
    assert bridge.has_unsaved_changes() is False
    assert bridge._confirm_close() is True
    assert window.close_dialogs == []


def test_dirty_native_close_cancels_or_confirms_and_lists_all_unsaved_apps(tmp_path):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    window = FakeWindow(confirmations=[False, True])
    bridge.set_window(window)
    assert bridge.set_unsaved_state("word", True)["success"] is True
    assert bridge.set_unsaved_state("excel", True)["has_unsaved_changes"] is True
    assert bridge._confirm_close() is False
    assert "Word belgesi" in window.close_dialogs[0][1]
    assert "Excel çalışma kitabı" in window.close_dialogs[0][1]
    assert bridge._confirm_close() is True
    assert bridge.has_unsaved_changes() is True  # A canceled close never discards state.
    bridge.set_unsaved_state("word", False)
    bridge.reset_current_file("excel")
    assert bridge.has_unsaved_changes() is False
    assert bridge._confirm_close() is True
    assert len(window.close_dialogs) == 2


def test_invalid_dirty_state_is_rejected_without_clearing_protection(tmp_path):
    bridge = desktop.NativeDesktopBridge(data_directory=tmp_path)
    bridge.set_unsaved_state("word", True)
    assert bridge.set_unsaved_state("word", "false")["success"] is False
    assert bridge.has_unsaved_changes() is True
    assert bridge._confirm_close() is False  # No available native window cannot confirm data loss.

from pathlib import Path
from docx import Document
from randomoffice.__main__ import main


def test_cli_reports_backend_status(capsys) -> None:
    assert main([]) == 0

    output = capsys.readouterr().out
    assert "RandomOffice backend is ready." in output
    assert ".docx, .xlsx, .pptx" in output


def test_cli_convert_command(tmp_path: Path, capsys) -> None:
    doc_path = tmp_path / "cli_test.docx"
    doc = Document()
    doc.add_paragraph("CLI Conversion Content")
    doc.save(str(doc_path))

    exit_code = main(["convert", str(doc_path), "--target", "text"])
    assert exit_code == 0
    output = capsys.readouterr().out
    assert "CLI Conversion Content" in output


def test_cli_parse_command(tmp_path: Path, capsys) -> None:
    doc_path = tmp_path / "cli_parse.docx"
    doc = Document()
    doc.add_paragraph("CLI Parsed Text")
    doc.save(str(doc_path))

    exit_code = main(["parse", str(doc_path)])
    assert exit_code == 0
    output = capsys.readouterr().out
    assert "CLI Parsed Text" in output


def test_desktop_cli_calls_native_entrypoint(monkeypatch) -> None:
    import sys
    from types import SimpleNamespace
    calls = []
    monkeypatch.setitem(sys.modules, "randomoffice.desktop", SimpleNamespace(run_desktop_app=lambda: calls.append(True)))
    assert main(["desktop"]) == 0
    assert calls == [True]

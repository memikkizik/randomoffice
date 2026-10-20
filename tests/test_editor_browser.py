"""Optional real-browser checks: uv run --extra browser pytest tests/test_editor_browser.py."""

import socket
import threading
import time

from docx import Document
from openpyxl import load_workbook
import pytest
import uvicorn

playwright = pytest.importorskip("playwright.sync_api")


@pytest.fixture(scope="module")
def browser_server():
    from randomoffice.api.app import create_app

    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        listener.listen(128)
        port = listener.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(create_app(), log_level="error"))
        thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
        thread.start()
        try:
            deadline = time.monotonic() + 10
            while not server.started:
                if not thread.is_alive() or time.monotonic() > deadline:
                    raise RuntimeError("Browser test server failed to start")
                time.sleep(0.02)
            with playwright.sync_playwright() as runner:
                try:
                    browser = runner.chromium.launch(channel="msedge", headless=True)
                except playwright.Error:
                    browser = runner.chromium.launch(headless=True)
                yield browser, f"http://127.0.0.1:{port}"
                browser.close()
        finally:
            server.should_exit = True
            thread.join(timeout=10)


@pytest.fixture
def page(browser_server):
    browser, url = browser_server
    context = browser.new_context(viewport={"width": 1280, "height": 850}, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(7000)
    page.goto(url)
    page.wait_for_function("typeof saveDocument === 'function'")
    yield page
    context.close()


def test_ribbon_tabs_have_no_runtime_errors(page):
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    for tab in ("insert", "view", "home"):
        page.locator(f"#tab-{tab}").click()
    assert not errors


def test_all_declared_ui_handlers_exist(page):
    missing = page.evaluate(r"""() => {
        const missing = new Set();
        document.querySelectorAll('*').forEach(element => {
            [...element.attributes].filter(attribute => attribute.name.startsWith('on')).forEach(attribute => {
                for (const match of attribute.value.matchAll(/(?<![.\w])([A-Za-z_$][\w$]*)\s*\(/g)) {
                    if (!['if','switch'].includes(match[1]) && typeof window[match[1]] !== 'function') missing.add(match[1]);
                }
            });
        });
        return [...missing];
    }""")
    assert missing == []


def test_excel_zero_and_wide_columns(page):
    page.evaluate("""ingestDocumentData({app:'excel', filename:'wide.xlsx', sheets:[
        {name:'Data',rows:[[0,...Array(25).fill(''),'last']]}]})""")
    assert page.locator('td[data-row="0"][data-col="0"]').inner_text() == "0"
    assert page.locator("#excel-sheet-grid tr").first.locator("th").last.inner_text() == "AA"


def test_formula_focus_and_save_preserve_expression(page, tmp_path):
    page.evaluate("""ingestDocumentData({app:'excel',filename:'formula.xlsx',sheets:[
        {name:'Data', rows:[[2,3,'=SUM(A1:B1)'],['','','']]}]})""")
    page.locator('td[data-row="0"][data-col="2"]').click()
    page.locator('td[data-row="1"][data-col="2"]').click()
    with page.expect_download() as download:
        page.evaluate("saveDocument()")
    target = tmp_path / "formula.xlsx"
    download.value.save_as(target)
    workbook = load_workbook(target)
    assert workbook.active["C1"].value == "=SUM(A1:B1)"
    workbook.close()


def test_word_open_does_not_duplicate_metadata_title(page):
    page.evaluate("""ingestDocumentData({app:'word',filename:'doc.docx',title:'Metadata',
        items:[{type:'paragraph',text:'Actual content',level:'p'}]})""")
    assert page.locator("#word-editor").inner_text().strip() == "Actual content"


def test_ppt_notes_and_tables_survive_ui_save(page, tmp_path):
    from pptx import Presentation

    page.evaluate("""ingestDocumentData({app:'powerpoint',filename:'deck.pptx',slides:[
        {title:'A slide',paragraphs:['Body'],notes:'Remember this',tables:[[['A','B']]]}]})""")
    with page.expect_download() as download:
        page.evaluate("saveDocument()")
    target = tmp_path / "deck.pptx"
    download.value.save_as(target)
    presentation = Presentation(target)
    slide = presentation.slides[0]
    assert slide.has_notes_slide
    assert "Remember this" in slide.notes_slide.notes_text_frame.text
    assert any(shape.has_table for shape in slide.shapes)


def test_blank_workbook_and_sheet_switch(page):
    page.evaluate("createBlankFromHub('excel')")
    assert page.locator("#excel-sheet-grid td").count() >= 300
    assert "Dizüstü" not in page.locator("#excel-sheet-grid").inner_text()
    page.evaluate("addNewSheetTab()")
    cell = page.locator('td[data-row="0"][data-col="0"]')
    cell.fill("Second sheet")
    cell.press("Tab")
    page.evaluate("switchExcelSheet(0)")
    assert page.locator('td[data-row="0"][data-col="0"]').inner_text() == ""


def test_storage_failure_does_not_break_typing(page):
    page.evaluate("() => { Storage.prototype.setItem = () => { throw new Error('Quota exceeded'); }; }")
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.locator("#word-editor").fill("Still editable")
    assert not errors
    assert "Kaydedildi" not in page.locator("#cloud-sync-status").inner_text()


def test_word_bold_survives_download(page, tmp_path):
    page.evaluate("document.getElementById('word-editor').innerHTML='<p><b>Bold text</b> plain</p>'")
    with page.expect_download() as download:
        page.evaluate("saveDocument()")
    target = tmp_path / "bold.docx"
    download.value.save_as(target)
    document = Document(target)
    assert any(run.bold and "Bold text" in run.text for p in document.paragraphs for run in p.runs)


@pytest.mark.parametrize(("formula", "expected"), [
    ("=SUM(A1:A3)", "5"), ("=AVERAGE(A1:A3)", "2,5"),
    ("=MIN(A1:A3)", "2"), ("=MAX(A1:A3)", "3"),
    ("=COUNT(A1:A3)", "2"), ("=SUM(A1:A3)*2+AA1", "14"),
    ("=1/0", "#DIV/0!"), ("=A4", "#CYCLE!"),
])
def test_formula_engine(page, formula, expected):
    page.evaluate("""ingestDocumentData({app:'excel',filename:'calc.xlsx',sheets:[
        {name:'Data',rows:[[2,...Array(25).fill(''),4],[3],[''],['=A4']]}]})""")
    assert page.evaluate("formula => evaluateFormulaDisplay(formula)", formula) == expected


def test_currency_keeps_number_and_export_format(page, tmp_path):
    page.evaluate("""ingestDocumentData({app:'excel',filename:'money.xlsx',sheets:[
        {name:'Data',rows:[[1234.5],['']]}]})""")
    page.locator('td[data-row="0"][data-col="0"]').click()
    page.evaluate("formatCurrencyExcel()")
    assert "1.234,50" in page.locator('td[data-row="0"][data-col="0"]').inner_text()
    with page.expect_download() as download:
        page.evaluate("saveDocument()")
    target = tmp_path / "money.xlsx"
    download.value.save_as(target)
    workbook = load_workbook(target)
    assert workbook.active["A1"].value == 1234.5
    assert "₺" in workbook.active["A1"].number_format
    workbook.close()


def test_toolbar_preserves_word_selection(page):
    page.evaluate("document.getElementById('word-editor').innerHTML='<p>Select this</p>'")
    page.locator("#word-editor").click()
    page.locator("#word-editor").press("Control+a")
    page.locator('button[title="Kalın (Ctrl+B)"]').click()
    assert page.locator("#word-editor b, #word-editor strong").inner_text() == "Select this"


def test_per_app_document_names_survive_reload(page):
    page.locator("#doc-name-editor").fill("My report.docx")
    page.locator("#pill-excel").click()
    page.locator("#doc-name-editor").fill("My budget.xlsx")
    page.locator("#pill-word").click()
    assert page.locator("#doc-name-editor").input_value() == "My report.docx"
    page.evaluate("persistState()")
    page.reload()
    assert page.locator("#doc-name-editor").input_value() == "My report.docx"
    page.locator("#pill-excel").click()
    assert page.locator("#doc-name-editor").input_value() == "My budget.xlsx"


def test_presenter_keyboard_and_notes_editing(page):
    page.locator("#pill-ppt").click()
    page.locator("#ppt-speaker-notes").fill("Private note")
    page.evaluate("createNewSlide()")
    page.evaluate("launchPresenterShow()")
    page.keyboard.press("ArrowRight")
    assert "2 / 2" in page.locator("#presenter-hud").inner_text()
    page.keyboard.press("ArrowLeft")
    assert "1 / 2" in page.locator("#presenter-hud").inner_text()
    page.keyboard.press("Escape")
    assert not page.locator("#presenter-modal").is_visible()
    page.evaluate("activeSlideIndex = 0; loadSlideToStage(0)")
    assert page.locator("#ppt-speaker-notes").input_value() == "Private note"


def test_compact_layout_keeps_status_bar_visible(page):
    page.set_viewport_size({"width": 960, "height": 640})
    bounds = page.locator(".desktop-status-bar").bounding_box()
    assert bounds["y"] + bounds["height"] <= 641
    assert page.locator("#word-workspace").bounding_box()["height"] > 200
    page.locator("#pill-ppt").click()
    card = page.locator("#slide-canvas-card").bounding_box()
    assert card["x"] >= 0
    assert card["x"] + card["width"] <= 961


def test_excel_undo_redo_updates_saved_model(page):
    page.evaluate("createBlankFromHub('excel')")
    cell = page.locator('td[data-row="0"][data-col="0"]')
    cell.fill("Changed")
    cell.press("Tab")
    page.evaluate("execOfficeAction('undo')")
    assert page.evaluate("excelModel.sheets[0].rows[0][0]") == ""
    page.evaluate("execOfficeAction('redo')")
    assert page.evaluate("excelModel.sheets[0].rows[0][0]") == "Changed"


def test_app_switch_survives_reload_without_editing(page):
    page.locator("#pill-excel").click()
    page.reload()
    assert page.locator("#excel-workspace").is_visible()


def test_deleting_later_sheet_keeps_current_sheet(page):
    page.evaluate("""ingestDocumentData({app:'excel',filename:'tabs.xlsx',sheets:[
        {name:'First',rows:[['first']]},{name:'Middle',rows:[['middle']]},{name:'Last',rows:[['last']]}]});
        switchExcelSheet(1); removeExcelSheet(2);""")
    assert page.locator('td[data-row="0"][data-col="0"]').inner_text() == "middle"


def test_native_save_rejection_is_reported_and_retry_works(page):
    page.evaluate("""() => { window.pywebview = {api: {
        save_native_file: async () => {throw new Error('Disk unavailable');}
    }}; }""")
    page.evaluate("saveDocument()")
    assert "Disk unavailable" in page.locator("#system-toast").inner_text()
    assert not page.locator('button[onclick="saveDocument()"]').first.is_disabled()
    page.evaluate("window.pywebview.api.save_native_file = async () => ({cancelled:true})")
    page.evaluate("saveDocument()")
    assert "iptal" in page.locator("#system-toast").inner_text()


def test_duplicate_save_only_calls_native_bridge_once(page):
    count = page.evaluate("""async () => {
        let calls=0, finish;
        window.pywebview={api:{save_native_file:()=>{calls++;return new Promise(resolve=>{finish=resolve;});}}};
        const first=saveDocument();
        await saveDocument();
        finish({cancelled:true});
        await first;
        return calls;
    }""")
    assert count == 1


def test_export_error_shows_server_message_without_download(page):
    page.route("**/api/v1/office/export?*", lambda route: route.fulfill(
        status=400, content_type="application/json", body='{"detail":"Invalid sheet name"}'
    ))
    page.evaluate("saveDocument()")
    assert "Invalid sheet name" in page.locator("#system-toast").inner_text()


def test_native_open_rejection_does_not_crash_page(page):
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.evaluate("""() => {window.pywebview={api:{open_native_file:async()=>{throw new Error('Dialog failed');}}};
        triggerOpenDialog();}""")
    assert "Dialog failed" in page.locator("#system-toast").inner_text()
    assert not errors

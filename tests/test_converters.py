from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from randomoffice.converters.contracts import ConversionTarget
from randomoffice.converters.pipeline import (
    ConversionPipeline,
    DocxConverter,
    PptxConverter,
    XlsxConverter,
)


def test_docx_converter_markdown_and_tables(tmp_path: Path) -> None:
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.core_properties.title = "Sample Document Title"
    doc.add_heading("Section 1", level=1)
    doc.add_paragraph("First paragraph text.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Val 1"
    table.rows[1].cells[1].text = "Val 2"

    doc.add_paragraph("Second paragraph text.")
    doc.save(str(doc_path))

    converter = DocxConverter()
    md_output = converter.convert(doc_path, ConversionTarget.MARKDOWN)

    assert "Sample Document Title" in md_output.content
    assert "# Section 1" in md_output.content
    assert "| Header A | Header B |" in md_output.content
    assert "| Val 1 | Val 2 |" in md_output.content
    assert md_output.metadata["title"] == "Sample Document Title"

    json_output = converter.convert(doc_path, ConversionTarget.JSON)
    assert '"type": "table"' in json_output.content
    assert '"Header A"' in json_output.content


def test_xlsx_converter_markdown_and_csv(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Q1", "Q2"])
    ws.append(["Widget", 100, 150])
    wb.save(str(xlsx_path))

    converter = XlsxConverter()
    md_output = converter.convert(xlsx_path, ConversionTarget.MARKDOWN)
    assert "## Sheet: Sales" in md_output.content
    assert "| Product | Q1 | Q2 |" in md_output.content
    assert "| Widget | 100 | 150 |" in md_output.content

    csv_output = converter.convert(xlsx_path, ConversionTarget.CSV)
    assert "Product,Q1,Q2" in csv_output.content
    assert "Widget,100,150" in csv_output.content


def test_pptx_converter_slides_and_tables(tmp_path: Path) -> None:
    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.core_properties.title = "Deck Title"
    
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Intro Slide"
    slide.placeholders[1].text = "Welcome to RandomOffice"

    prs.save(str(pptx_path))

    converter = PptxConverter()
    md_output = converter.convert(pptx_path, ConversionTarget.MARKDOWN)
    assert "## Intro Slide (Slide 1)" in md_output.content
    assert "Welcome to RandomOffice" in md_output.content


def test_conversion_pipeline_with_caching(tmp_path: Path) -> None:
    doc_path = tmp_path / "cached_sample.docx"
    doc = Document()
    doc.add_paragraph("Hello Cache")
    doc.save(str(doc_path))

    pipeline = ConversionPipeline()
    res1 = pipeline.convert(doc_path, ConversionTarget.TEXT, use_cache=True)
    assert "Hello Cache" in res1.content

    # Second call should fetch from cache
    res2 = pipeline.convert(doc_path, ConversionTarget.TEXT, use_cache=True)
    assert res1 is res2
